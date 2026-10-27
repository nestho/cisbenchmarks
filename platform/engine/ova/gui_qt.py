from __future__ import annotations


from datetime import datetime
import os
from pathlib import Path
import sys
import threading

from PySide6.QtCore import QObject, QProcess, QSettings, Qt, QTimer, Signal
from PySide6.QtGui import QColor
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QFrame,
    QGraphicsDropShadowEffect,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QLayout,
    QMainWindow,
    QMessageBox,
    QPlainTextEdit,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QSpinBox,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

try:
    from .feed_manager import (
        OperationCancelled,
        all_feed_specs,
        bootstrap_all,
        feed_freshness,
        summarize_feed_freshness,
    )
    from .collector_commands import collector_command, collector_label
    from .coverage_registry import coverage_rows, scope_preview
    from .ssh_collector import (
        CollectResult,
        JumpHost,
        SSHOptions,
        Target,
        collect_many,
        export_collector_bash,
        parse_target_list,
    )
    from . import __version__
except ImportError:  # launched as a script: python offline_va/gui_qt.py
    import sys as _sys
    from pathlib import Path as _Path
    _root = _Path(__file__).resolve().parent.parent
    if str(_root) not in _sys.path:
        _sys.path.insert(0, str(_root))
    from offline_va.feed_manager import (
        OperationCancelled,
        all_feed_specs,
        bootstrap_all,
        feed_freshness,
        summarize_feed_freshness,
    )
    from offline_va.collector_commands import collector_command, collector_label
    from offline_va.coverage_registry import coverage_rows, scope_preview
    from offline_va.ssh_collector import (
        CollectResult,
        JumpHost,
        SSHOptions,
        Target,
        collect_many,
        export_collector_bash,
        parse_target_list,
    )
    from offline_va import __version__


PRODUCT_NAME = "OfflineVA"
PRODUCT_VERSION = __version__


SUPPORTED_OS_FAMILIES = (
    ("auto", "Auto-detect from collector", "Use the OS identity recorded by the collector."),
    ("redhat", "Red Hat Enterprise Linux", "Evaluate Red Hat OVAL (6-9) and CSAF (10+) vendor evidence."),
    ("ubuntu", "Ubuntu", "Evaluate Canonical USN/CVE OVAL (dpkginfo) for supported codenames."),
    ("suse", "SUSE Linux Enterprise", "Evaluate SUSE vendor evidence only."),
    ("debian", "Debian GNU/Linux", "Evaluate Debian Security Tracker evidence only."),
    ("huawei", "Huawei Cloud EulerOS", "Evaluate Huawei vendor evidence only."),
)

LIMITED_OS_FAMILIES = (
    ("oracle", "Oracle Linux (coverage limited)"),
    ("rocky", "Rocky Linux (coverage limited)"),
    ("alma", "AlmaLinux (coverage limited)"),
    ("centos", "CentOS (coverage limited)"),
    ("amazon", "Amazon Linux (coverage limited)"),
    ("fedora", "Fedora (coverage limited)"),
    ("opensuse", "openSUSE (coverage limited)"),
    ("alpine", "Alpine Linux (coverage limited)"),
    ("arch", "Arch Linux (coverage limited)"),
    ("mint", "Linux Mint (coverage limited)"),
    ("kali", "Kali Linux (coverage limited)"),
    ("openeuler", "openEuler (coverage limited)"),
)


def system_uses_dark_theme() -> bool:
    try:
        import winreg

        path = r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize"
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, path) as key:
            value, _ = winreg.QueryValueEx(key, "AppsUseLightTheme")
        return int(value) == 0
    except (ImportError, OSError, TypeError, ValueError):
        return False


class WorkerSignals(QObject):
    log = Signal(str)
    feed_progress = Signal(float, str, str)
    feed_done = Signal(object, object)
    operation_error = Signal(str, str)
    stopped = Signal(str)


class Toast(QFrame):
    def __init__(self, parent: QWidget, title: str, message: str, kind: str = "success"):
        super().__init__(parent)
        self.setObjectName("toast")
        self.setProperty("kind", kind)
        self.setFixedWidth(410)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(14, 12, 14, 12)
        dot = QLabel("●")
        dot.setObjectName("toastDot")
        text = QVBoxLayout()
        title_label = QLabel(title)
        title_label.setObjectName("toastTitle")
        message_label = QLabel(message)
        message_label.setObjectName("toastMessage")
        message_label.setWordWrap(True)
        text.addWidget(title_label)
        text.addWidget(message_label)
        close = QPushButton("×")
        close.setObjectName("toastClose")
        close.setFixedSize(26, 26)
        close.clicked.connect(self.close)
        layout.addWidget(dot, 0, Qt.AlignmentFlag.AlignTop)
        layout.addLayout(text, 1)
        layout.addWidget(close, 0, Qt.AlignmentFlag.AlignTop)
        shadow = QGraphicsDropShadowEffect(self)
        shadow.setBlurRadius(28)
        shadow.setOffset(0, 8)
        shadow.setColor(QColor(0, 0, 0, 90))
        self.setGraphicsEffect(shadow)
        self.adjustSize()


class StatCard(QFrame):
    def __init__(self, label: str, color: str, initial: str = "0"):
        super().__init__()
        self.setProperty("statCard", True)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 16, 18, 14)
        layout.setSpacing(6)
        title = QLabel(label.upper())
        title.setObjectName("statLabel")
        self.value = QLabel(initial)
        self.value.setObjectName("statValue")
        accent = QFrame()
        accent.setFixedHeight(3)
        accent.setStyleSheet(
            f"background: {color}; border-radius: 2px; margin-top: 4px;"
        )
        layout.addWidget(title)
        layout.addWidget(self.value)
        layout.addStretch(1)
        layout.addWidget(accent)


class OfflineVAWindow(QMainWindow):
    """OfflineVA UI using the same professional design language as OfflineCompliance."""

    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"{PRODUCT_NAME} — {PRODUCT_VERSION}")
        self.resize(1480, 920)
        self.setMinimumSize(1180, 760)

        self.project_dir = Path(__file__).resolve().parent.parent
        self.dark_theme = system_uses_dark_theme()
        self.cancel_event = threading.Event()
        self.process: QProcess | None = None
        self.is_busy = False
        self.active_toast: Toast | None = None
        self.last_report: Path | None = None
        self.last_full_report: Path | None = None
        self.feed_rows: list[dict] = []
        self.coverage_inventory = None
        self.signals = WorkerSignals()
        self.signals.log.connect(self._append_log)
        self.signals.feed_progress.connect(self._on_feed_progress)
        self.signals.feed_done.connect(self._on_feed_done)
        self.signals.operation_error.connect(self._on_operation_error)
        self.signals.stopped.connect(self._on_stopped)

        self._build_ui()
        self._apply_theme()
        self._show_page(0)
        QTimer.singleShot(150, self._refresh_feed_health)
        self.theme_timer = QTimer(self)
        self.theme_timer.timeout.connect(self._monitor_system_theme)
        self.theme_timer.start(2000)

    # ---------- composition ----------
    def _build_ui(self):
        root = QWidget()
        root_layout = QHBoxLayout(root)
        root_layout.setContentsMargins(0, 0, 0, 0)
        root_layout.setSpacing(0)
        self.setCentralWidget(root)

        self.sidebar = QFrame()
        self.sidebar.setObjectName("sidebar")
        self.sidebar.setFixedWidth(248)
        side = QVBoxLayout(self.sidebar)
        side.setContentsMargins(18, 24, 18, 18)
        side.setSpacing(4)

        brand = QHBoxLayout()
        brand.setSpacing(12)
        mark = QLabel("VA")
        mark.setObjectName("brandMark")
        mark.setAlignment(Qt.AlignmentFlag.AlignCenter)
        mark.setFixedSize(40, 40)
        name_box = QVBoxLayout()
        name_box.setSpacing(1)
        name = QLabel("OfflineVA")
        name.setObjectName("brandName")
        edition = QLabel("Vendor evidence engine")
        edition.setObjectName("brandEdition")
        name_box.addWidget(name)
        name_box.addWidget(edition)
        brand.addWidget(mark)
        brand.addLayout(name_box, 1)
        side.addLayout(brand)
        side.addSpacing(28)

        nav_hint = QLabel("WORKSPACE")
        nav_hint.setObjectName("navSection")
        side.addWidget(nav_hint)
        side.addSpacing(6)

        self.nav_buttons: list[QPushButton] = []
        for index, (label, number) in enumerate((
            ("Overview", "01"),
            ("Audit", "02"),
            ("Remote", "03"),
            ("Feeds", "04"),
            ("Reports", "05"),
            ("Settings", "06"),
        )):
            button = QPushButton(f"  {number}   {label}")
            button.setProperty("nav", True)
            button.setCheckable(True)
            button.setCursor(Qt.CursorShape.PointingHandCursor)
            button.setMinimumHeight(40)
            button.clicked.connect(lambda _checked=False, page=index: self._show_page(page))
            self.nav_buttons.append(button)
            side.addWidget(button)
        side.addStretch(1)

        safety = QFrame()
        safety.setObjectName("safetyCard")
        safety_layout = QVBoxLayout(safety)
        safety_layout.setContentsMargins(14, 12, 14, 12)
        safety_layout.setSpacing(4)
        safe_title = QLabel("●  VENDOR-FIRST")
        safe_title.setObjectName("safeTitle")
        safe_text = QLabel("NVD enriches findings only.\nUnknown scope is never confirmed.")
        safe_text.setObjectName("safeText")
        safety_layout.addWidget(safe_title)
        safety_layout.addWidget(safe_text)
        side.addWidget(safety)
        side.addSpacing(8)
        version = QLabel(f"v{PRODUCT_VERSION}")
        version.setObjectName("sideVersion")
        version.setAlignment(Qt.AlignmentFlag.AlignCenter)
        side.addWidget(version)

        workspace = QFrame()
        workspace.setObjectName("workspace")
        work_layout = QVBoxLayout(workspace)
        work_layout.setContentsMargins(32, 24, 32, 18)
        work_layout.setSpacing(18)

        top = QHBoxLayout()
        top.setSpacing(16)
        titles = QVBoxLayout()
        titles.setSpacing(4)
        self.page_title = QLabel("Overview")
        self.page_title.setObjectName("pageTitle")
        self.page_subtitle = QLabel("Host, feed and report readiness at a glance")
        self.page_subtitle.setObjectName("pageSubtitle")
        titles.addWidget(self.page_title)
        titles.addWidget(self.page_subtitle)
        self.feed_pill = QLabel("●  Checking feed snapshot")
        self.feed_pill.setObjectName("warningPill")
        self.feed_pill.setAlignment(Qt.AlignmentFlag.AlignVCenter)
        top.addLayout(titles, 1)
        top.addWidget(self.feed_pill, 0, Qt.AlignmentFlag.AlignTop)
        work_layout.addLayout(top)

        self.stack = QStackedWidget()
        self.stack.addWidget(self._build_overview_page())
        self.stack.addWidget(self._build_audit_page())
        self.stack.addWidget(self._build_remote_page())
        self.stack.addWidget(self._build_feeds_page())
        self.stack.addWidget(self._build_reports_page())
        self.stack.addWidget(self._build_settings_page())
        work_layout.addWidget(self.stack, 1)

        footer = QHBoxLayout()
        self.status_label = QLabel("Ready")
        self.status_label.setObjectName("footerText")
        self.theme_label = QLabel()
        self.theme_label.setObjectName("footerText")
        footer.addWidget(self.status_label)
        footer.addStretch(1)
        footer.addWidget(self.theme_label)
        work_layout.addLayout(footer)

        root_layout.addWidget(self.sidebar)
        root_layout.addWidget(workspace, 1)

    def _page(self) -> tuple[QWidget, QVBoxLayout]:
        page = QScrollArea()
        page.setWidgetResizable(True)
        page.setFrameShape(QFrame.Shape.NoFrame)
        page.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        content = QWidget()
        content.setObjectName("pageContent")
        layout = QVBoxLayout(content)
        layout.setSizeConstraint(QLayout.SizeConstraint.SetMinimumSize)
        layout.setContentsMargins(0, 0, 4, 8)
        layout.setSpacing(16)
        page.setWidget(content)
        return page, layout

    def _card(self, title: str, subtitle: str = "") -> tuple[QFrame, QVBoxLayout]:
        card = QFrame()
        card.setProperty("card", True)
        layout = QVBoxLayout(card)
        layout.setContentsMargins(22, 18, 22, 18)
        layout.setSpacing(12)
        heading = QLabel(title)
        heading.setObjectName("cardTitle")
        layout.addWidget(heading)
        if subtitle:
            note = QLabel(subtitle)
            note.setObjectName("cardSubtitle")
            note.setWordWrap(True)
            layout.addWidget(note)
        return card, layout

    def _build_overview_page(self):
        page, layout = self._page()
        hero, body = self._card(
            "Advisory-level vulnerability analysis",
            "One host/advisory result groups every matched package and CVE. Confirmed totals exclude unresolved scope candidates.",
        )
        hero.setObjectName("heroCard")
        actions = QHBoxLayout()
        run = QPushButton("Configure & run audit")
        run.setProperty("primary", True)
        run.clicked.connect(lambda: self._show_page(1))
        feeds = QPushButton("Review feed snapshot")
        feeds.setProperty("secondary", True)
        feeds.clicked.connect(lambda: self._show_page(3))
        actions.addWidget(run)
        actions.addWidget(feeds)
        actions.addStretch(1)
        body.addLayout(actions)
        layout.addWidget(hero)

        readiness = QGridLayout()
        readiness.setHorizontalSpacing(12)
        input_card, input_body = self._card("Inventory input")
        self.overview_input = QLabel("No collector TXT selected")
        self.overview_input.setObjectName("readinessValue")
        self.overview_input.setWordWrap(True)
        input_body.addWidget(self.overview_input)
        input_action = QPushButton("Select inventory")
        input_action.setProperty("link", True)
        input_action.clicked.connect(lambda: self._show_page(1))
        input_body.addWidget(input_action, 0, Qt.AlignmentFlag.AlignLeft)

        feed_card, feed_body = self._card("Feed applicability")
        self.overview_feed = QLabel("Checking local snapshot")
        self.overview_feed.setObjectName("readinessValue")
        self.overview_feed.setWordWrap(True)
        feed_body.addWidget(self.overview_feed)
        feed_action = QPushButton("Open feed manager")
        feed_action.setProperty("link", True)
        feed_action.clicked.connect(lambda: self._show_page(3))
        feed_body.addWidget(feed_action, 0, Qt.AlignmentFlag.AlignLeft)

        report_card, report_body = self._card("Report destination")
        self.overview_output = QLabel("No CVE / full assessment generated in this session")
        self.overview_output.setObjectName("readinessValue")
        self.overview_output.setWordWrap(True)
        report_body.addWidget(self.overview_output)
        report_action = QPushButton("Open reports")
        report_action.setProperty("link", True)
        report_action.clicked.connect(lambda: self._show_page(4))
        report_body.addWidget(report_action, 0, Qt.AlignmentFlag.AlignLeft)
        readiness.addWidget(input_card, 0, 0)
        readiness.addWidget(feed_card, 0, 1)
        readiness.addWidget(report_card, 0, 2)
        layout.addLayout(readiness)

        snapshot, snap_body = self._card("Latest report snapshot", "Confirmed severity counts are Nessus-style advisory findings; review candidates remain separate.")
        stats = QHBoxLayout()
        self.report_stats = {}
        for label, color in (
            ("Confirmed", "#3B82F6"), ("Critical", "#91243E"),
            ("High", "#DD4B50"), ("Medium", "#F18C43"),
            ("Low", "#F8C851"), ("Review", "#8B5CF6"),
        ):
            card = StatCard(label, color)
            self.report_stats[label] = card.value
            stats.addWidget(card)
        snap_body.addLayout(stats)
        layout.addWidget(snapshot)
        layout.addStretch(1)
        return page

    def _build_audit_page(self):
        page, layout = self._page()
        workspace, body = self._card(
            "Audit workspace",
            "Collector TXT (LINUX_REPORT_SAFE_V6 or OfflineVA inventory) or a folder of hosts. "
            "Each host produces CVE workbook + full OS Security Assessment (enterprise template).",
        )

        scope_row = QHBoxLayout()
        scope_row.setSpacing(12)
        scope_label = QLabel("Target operating system / سیستم‌عامل هدف")
        scope_label.setObjectName("fieldLabel")
        self.os_family_combo = QComboBox()
        self.os_family_combo.setObjectName("targetOsFamilyCombo")
        self.os_family_combo.setMinimumWidth(265)
        for value, label, tooltip in SUPPORTED_OS_FAMILIES:
            self.os_family_combo.addItem(label, value)
            self.os_family_combo.setItemData(
                self.os_family_combo.count() - 1,
                tooltip,
                Qt.ItemDataRole.ToolTipRole,
            )
        self.os_family_combo.insertSeparator(self.os_family_combo.count())
        for value, label in LIMITED_OS_FAMILIES:
            self.os_family_combo.addItem(label, value)
            self.os_family_combo.setItemData(
                self.os_family_combo.count() - 1,
                "The OS identity is recorded, but full vendor-first vulnerability coverage is not available yet.",
                Qt.ItemDataRole.ToolTipRole,
            )
        saved_family = str(QSettings("OfflineVA", f"{PRODUCT_NAME}-3.6").value("audit/os_family", "auto"))
        saved_index = self.os_family_combo.findData(saved_family)
        self.os_family_combo.setCurrentIndex(saved_index if saved_index >= 0 else 0)
        self.os_family_combo.currentIndexChanged.connect(self._on_os_family_changed)
        self.os_scope_hint = QLabel()
        self.os_scope_hint.setObjectName("inlineStatus")
        self.os_scope_hint.setWordWrap(True)
        scope_row.addWidget(scope_label)
        scope_row.addWidget(self.os_family_combo, 1)
        body.addLayout(scope_row)
        body.addWidget(self.os_scope_hint)
        self._on_os_family_changed()
        self.coverage_hint = QLabel(scope_preview(self._selected_os_family(), self.project_dir / "feeds"))
        self.coverage_hint.setWordWrap(True)
        self.coverage_hint.setObjectName("inlineStatus")
        body.addWidget(self.coverage_hint)

        command_row = QHBoxLayout()
        self.collector_toggle = QCheckBox(collector_label(self._selected_os_family()))
        self.collector_toggle.setChecked(True)
        self.copy_collector_button = QPushButton("Copy full command / کپی دستور")
        self.copy_collector_button.setProperty("secondary", True)
        self.copy_collector_button.clicked.connect(self._copy_collector_command)
        command_row.addWidget(self.collector_toggle)
        command_row.addStretch(1)
        command_row.addWidget(self.copy_collector_button)
        body.addLayout(command_row)
        self.collector_edit = QPlainTextEdit()
        self.collector_edit.setReadOnly(True)
        self.collector_edit.setObjectName("monoPanel")
        self.collector_edit.setFixedHeight(126)
        self.collector_edit.setPlainText(collector_command(self._selected_os_family()))
        self.collector_toggle.toggled.connect(self.collector_edit.setVisible)
        body.addWidget(self.collector_edit)
        self.collector_help = QLabel("دستور مخصوص انتخاب بالا را روی Linux هدف اجرا و تمام خروجی را در TXT ذخیره کنید. بدون hostname -f، نصب یا refresh مخزن؛ URLها پنهان می‌شوند. تنظیمات مخزن و امضا فقط metadata هستند، نه تأیید وضعیت مؤثر یا سلامت پکیج. گزینه‌های coverage limited پوشش کامل VA ندارند.")
        self.collector_help.setWordWrap(True)
        self.collector_help.setObjectName("inlineStatus")
        body.addWidget(self.collector_help)

        grid = QGridLayout()
        grid.setHorizontalSpacing(10)
        grid.setVerticalSpacing(9)
        self.input_edit = QLineEdit(str(self.project_dir / "input"))
        self.feeds_edit = QLineEdit(str(self.project_dir / "feeds"))
        self.feeds_edit.editingFinished.connect(self._refresh_feed_health)
        self.output_edit = QLineEdit(str(self.project_dir / "output"))
        self.context_edit = QLineEdit()
        self.context_edit.setPlaceholderText("Optional: full collector for the same host when input has only name + version")
        self.input_edit.textChanged.connect(self._sync_overview)
        self.input_edit.editingFinished.connect(self._suggest_metadata_context)
        self.output_edit.textChanged.connect(self._sync_overview)

        def add_path(row: int, label: str, edit: QLineEdit, buttons: list[tuple[str, object]]):
            caption = QLabel(label)
            caption.setObjectName("fieldLabel")
            grid.addWidget(caption, row * 2, 0, 1, 4)
            edit.setMinimumHeight(34)
            grid.addWidget(edit, row * 2 + 1, 0, 1, 3)
            box = QHBoxLayout()
            box.setContentsMargins(0, 0, 0, 0)
            box.setSpacing(6)
            for text, slot in buttons:
                button = QPushButton(text)
                button.setProperty("secondary", True)
                button.clicked.connect(slot)
                box.addWidget(button)
            holder = QWidget()
            holder.setLayout(box)
            grid.addWidget(holder, row * 2 + 1, 3)
            return caption, edit, holder

        add_path(0, "Collector inventory TXT / folder", self.input_edit, [("TXT", self._browse_input_file), ("Folder", self._browse_input_folder)])
        add_path(1, "Feed repository", self.feeds_edit, [("Browse", self._browse_feeds)])
        add_path(2, "Output folder (CVE + full assessment)", self.output_edit, [("XLSX", self._browse_output_file), ("Folder", self._browse_output_folder)])
        self.context_widgets = add_path(3, "Full collector from this same host / خروجی کامل همان میزبان", self.context_edit, [("TXT", self._browse_context)])
        self.context_toggle = QCheckBox("Package-only input / ورودی من فقط نام و نسخهٔ پکیج است")
        self.context_toggle.setToolTip("فقط برای فهرست دو‌ستونه: خروجی کامل همان میزبان را پیوست کنید. اگر دستور بالا را اجرا کرده‌اید، نیازی به این گزینه نیست.")
        self.context_toggle.toggled.connect(self._toggle_metadata_context)
        body.addWidget(self.context_toggle)
        self._toggle_metadata_context(False)
        grid.setColumnStretch(0, 1)
        grid.setColumnStretch(1, 1)
        grid.setColumnStretch(2, 1)
        body.addLayout(grid)

        options = QHBoxLayout()
        self.enrich_check = QCheckBox("NVD + KEV + EPSS enrichment")
        self.enrich_check.setChecked(True)
        self.auto_feeds_check = QCheckBox("Update feeds before every audit")
        self.full_assessment_check = QCheckBox("Full OS Security Assessment Excel (enterprise template)")
        self.full_assessment_check.setChecked(True)
        self.full_assessment_check.setToolTip(
            "After CVE analysis, also write *_OS_Security_Assessment.xlsx with Host/Package/Users/… "
            "and Vulnerabilities mapped from OfflineVA findings (same design as the CT assessment template)."
        )
        options.addWidget(self.enrich_check)
        options.addWidget(self.auto_feeds_check)
        options.addWidget(self.full_assessment_check)
        options.addStretch(1)
        body.addLayout(options)

        actions = QHBoxLayout()
        self.audit_ready = QLabel("Ready to validate paths and feed freshness")
        self.audit_ready.setObjectName("inlineStatus")
        self.audit_ready.setWordWrap(True)
        self.stop_button = QPushButton("Stop")
        self.stop_button.setProperty("danger", True)
        self.stop_button.setEnabled(False)
        self.stop_button.clicked.connect(self._stop_operation)
        self.run_button = QPushButton("Run full offline audit")
        self.run_button.setProperty("success", True)
        self.run_button.clicked.connect(self._request_audit)
        actions.addWidget(self.audit_ready, 1)
        actions.addWidget(self.stop_button)
        actions.addWidget(self.run_button)
        body.addLayout(actions)
        self.progress = QProgressBar()
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        self.progress.setTextVisible(False)
        body.addWidget(self.progress)
        layout.addWidget(workspace)

        log_card, log_body = self._card("Execution log", "Feed and audit output is streamed without blocking the interface.")
        self.log_edit = QPlainTextEdit()
        self.log_edit.setReadOnly(True)
        self.log_edit.setMinimumHeight(120)
        self.log_edit.setObjectName("monoPanel")
        self.log_edit.setPlaceholderText("Execution messages will appear here.")
        log_body.addWidget(self.log_edit)
        layout.addWidget(log_card, 1)
        return page

    def _build_remote_page(self):
        """SSH remote collection: generate safe bash + pull inventory TXT from hosts."""
        page, layout = self._page()
        intro, intro_body = self._card(
            "Remote inventory over SSH",
            "Runs only the fixed read-only collector script. Supports password, key, jump host, "
            "and legacy algorithms for older OpenSSH / dropbear. Authorized use only.",
        )
        help_lbl = QLabel(
            "One host or a list. Output TXT files land in the chosen folder and can be fed straight into Audit. "
            "Nothing is written on the remote host; the script is piped to `sh -s`."
        )
        help_lbl.setWordWrap(True)
        help_lbl.setObjectName("inlineStatus")
        intro_body.addWidget(help_lbl)
        layout.addWidget(intro)

        # --- targets ---
        targets_card, targets_body = self._card(
            "Targets",
            "One per line: host | user@host | user@host:port | user@host:port,os_family,label   (# comments ok)",
        )
        self.remote_targets_edit = QPlainTextEdit()
        self.remote_targets_edit.setPlaceholderText(
            "root@10.0.0.12\n"
            "admin@legacy-box.local:22,redhat,dc1-app\n"
            "oracle@192.168.1.50:2222,oracle\n"
            "# jump will be applied to all lines below when enabled"
        )
        self.remote_targets_edit.setObjectName("monoPanel")
        self.remote_targets_edit.setMinimumHeight(110)
        targets_body.addWidget(self.remote_targets_edit)

        target_row = QHBoxLayout()
        target_row.addWidget(QLabel("Default user"))
        self.remote_default_user = QLineEdit("root")
        self.remote_default_user.setFixedWidth(120)
        target_row.addWidget(self.remote_default_user)
        target_row.addWidget(QLabel("Default port"))
        self.remote_default_port = QSpinBox()
        self.remote_default_port.setRange(1, 65535)
        self.remote_default_port.setValue(22)
        self.remote_default_port.setFixedWidth(90)
        target_row.addWidget(self.remote_default_port)
        target_row.addWidget(QLabel("Default OS family"))
        self.remote_default_family = QComboBox()
        for key, label, _ in SUPPORTED_OS_FAMILIES:
            self.remote_default_family.addItem(label, key)
        for key, label in LIMITED_OS_FAMILIES:
            self.remote_default_family.addItem(label, key)
        self.remote_default_family.setCurrentIndex(0)
        target_row.addWidget(self.remote_default_family, 1)
        targets_body.addLayout(target_row)
        layout.addWidget(targets_card)

        # --- auth ---
        auth_card, auth_body = self._card("Authentication", "Same credentials apply to every target unless you override per line later.")
        auth_grid = QGridLayout()
        auth_grid.addWidget(QLabel("Auth method"), 0, 0)
        self.remote_auth_method = QComboBox()
        self.remote_auth_method.addItem("Password", "password")
        self.remote_auth_method.addItem("Private key", "key")
        self.remote_auth_method.addItem("Key + passphrase", "key_pass")
        auth_grid.addWidget(self.remote_auth_method, 0, 1)

        auth_grid.addWidget(QLabel("Password"), 1, 0)
        self.remote_password = QLineEdit()
        self.remote_password.setEchoMode(QLineEdit.EchoMode.Password)
        self.remote_password.setPlaceholderText("Leave empty if using a key")
        auth_grid.addWidget(self.remote_password, 1, 1)

        auth_grid.addWidget(QLabel("Private key path"), 2, 0)
        key_row = QHBoxLayout()
        self.remote_key_path = QLineEdit()
        self.remote_key_path.setPlaceholderText("e.g. C:\\Users\\you\\.ssh\\id_rsa  or  /home/you/.ssh/id_rsa")
        key_browse = QPushButton("Browse")
        key_browse.setProperty("secondary", True)
        key_browse.clicked.connect(self._remote_browse_key)
        key_row.addWidget(self.remote_key_path, 1)
        key_row.addWidget(key_browse)
        auth_grid.addLayout(key_row, 2, 1)

        auth_grid.addWidget(QLabel("Key passphrase"), 3, 0)
        self.remote_key_pass = QLineEdit()
        self.remote_key_pass.setEchoMode(QLineEdit.EchoMode.Password)
        auth_grid.addWidget(self.remote_key_pass, 3, 1)
        auth_body.addLayout(auth_grid)
        layout.addWidget(auth_card)

        # --- jump host ---
        jump_card, jump_body = self._card("Jump / bastion host (optional)")
        self.remote_use_jump = QCheckBox("Connect through a jump host")
        jump_body.addWidget(self.remote_use_jump)
        jump_grid = QGridLayout()
        jump_grid.addWidget(QLabel("Jump host"), 0, 0)
        self.remote_jump_host = QLineEdit()
        self.remote_jump_host.setPlaceholderText("bastion.example.com")
        jump_grid.addWidget(self.remote_jump_host, 0, 1)
        jump_grid.addWidget(QLabel("Port"), 0, 2)
        self.remote_jump_port = QSpinBox()
        self.remote_jump_port.setRange(1, 65535)
        self.remote_jump_port.setValue(22)
        jump_grid.addWidget(self.remote_jump_port, 0, 3)
        jump_grid.addWidget(QLabel("User"), 1, 0)
        self.remote_jump_user = QLineEdit("root")
        jump_grid.addWidget(self.remote_jump_user, 1, 1)
        jump_grid.addWidget(QLabel("Password"), 1, 2)
        self.remote_jump_password = QLineEdit()
        self.remote_jump_password.setEchoMode(QLineEdit.EchoMode.Password)
        jump_grid.addWidget(self.remote_jump_password, 1, 3)
        jump_grid.addWidget(QLabel("Key path"), 2, 0)
        self.remote_jump_key = QLineEdit()
        jump_grid.addWidget(self.remote_jump_key, 2, 1, 1, 3)
        jump_body.addLayout(jump_grid)
        layout.addWidget(jump_card)

        # --- options + output ---
        opts_card, opts_body = self._card("Options & output")
        opts_row = QHBoxLayout()
        self.remote_legacy = QCheckBox("Allow legacy SSH algorithms (old servers)")
        self.remote_legacy.setChecked(True)
        opts_row.addWidget(self.remote_legacy)
        self.remote_auto_hostkey = QCheckBox("Auto-accept unknown host keys")
        self.remote_auto_hostkey.setChecked(True)
        opts_row.addWidget(self.remote_auto_hostkey)
        opts_row.addWidget(QLabel("Workers"))
        self.remote_workers = QSpinBox()
        self.remote_workers.setRange(1, 8)
        self.remote_workers.setValue(1)
        self.remote_workers.setFixedWidth(60)
        opts_row.addWidget(self.remote_workers)
        opts_row.addWidget(QLabel("Cmd timeout (s)"))
        self.remote_cmd_timeout = QSpinBox()
        self.remote_cmd_timeout.setRange(30, 3600)
        self.remote_cmd_timeout.setValue(300)
        self.remote_cmd_timeout.setFixedWidth(80)
        opts_row.addWidget(self.remote_cmd_timeout)
        opts_row.addStretch(1)
        opts_body.addLayout(opts_row)

        out_row = QHBoxLayout()
        out_row.addWidget(QLabel("Output folder"))
        self.remote_output_dir = QLineEdit()
        default_out = str(Path.home() / "OfflineVA" / "inventories")
        self.remote_output_dir.setText(default_out)
        out_row.addWidget(self.remote_output_dir, 1)
        out_browse = QPushButton("Browse")
        out_browse.setProperty("secondary", True)
        out_browse.clicked.connect(self._remote_browse_output)
        out_row.addWidget(out_browse)
        opts_body.addLayout(out_row)
        layout.addWidget(opts_card)

        # --- actions ---
        actions_card, actions_body = self._card("Actions")
        btn_row = QHBoxLayout()
        self.remote_export_bash = QPushButton("Export collector .sh")
        self.remote_export_bash.setProperty("secondary", True)
        self.remote_export_bash.clicked.connect(self._remote_export_bash)
        self.remote_copy_bash = QPushButton("Copy collector bash")
        self.remote_copy_bash.setProperty("secondary", True)
        self.remote_copy_bash.clicked.connect(self._remote_copy_bash)
        self.remote_stop_btn = QPushButton("Stop")
        self.remote_stop_btn.setProperty("danger", True)
        self.remote_stop_btn.setEnabled(False)
        self.remote_stop_btn.clicked.connect(self._remote_stop)
        self.remote_run_btn = QPushButton("Collect via SSH")
        self.remote_run_btn.setProperty("success", True)
        self.remote_run_btn.clicked.connect(self._remote_start)
        btn_row.addWidget(self.remote_export_bash)
        btn_row.addWidget(self.remote_copy_bash)
        btn_row.addStretch(1)
        btn_row.addWidget(self.remote_stop_btn)
        btn_row.addWidget(self.remote_run_btn)
        actions_body.addLayout(btn_row)
        self.remote_progress = QProgressBar()
        self.remote_progress.setRange(0, 100)
        self.remote_progress.setValue(0)
        actions_body.addWidget(self.remote_progress)
        self.remote_status = QLabel("Idle")
        self.remote_status.setObjectName("inlineStatus")
        actions_body.addWidget(self.remote_status)
        layout.addWidget(actions_card)

        log_card, log_body = self._card("Remote collection log")
        self.remote_log = QPlainTextEdit()
        self.remote_log.setReadOnly(True)
        self.remote_log.setObjectName("monoPanel")
        self.remote_log.setMinimumHeight(140)
        self.remote_log.setPlaceholderText("SSH connection and collection messages appear here.")
        log_body.addWidget(self.remote_log)
        layout.addWidget(log_card, 1)

        self._remote_stop_event = threading.Event()
        self._remote_thread: threading.Thread | None = None
        return page

    def _build_feeds_page(self):
        page, layout = self._page()
        coverage_card, coverage_body = self._card(
            "OS coverage / پوشش سیستم‌عامل",
            "Implementation catalog, not a security result. Fresh feeds do not mean all OS versions are covered. Architecture and support stream must be evaluated per host.",
        )
        self.coverage_table = QTableWidget(0, 6)
        self.coverage_table.setHorizontalHeaderLabels(("Distribution", "Implemented releases", "Architecture", "Assessment capability", "Vendor feeds C/S/M/I", "Scope limits"))
        self.coverage_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.coverage_table.verticalHeader().setVisible(False)
        self.coverage_table.setMinimumHeight(230)
        self.coverage_table.setWordWrap(True)
        self.coverage_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        coverage_body.addWidget(self.coverage_table)
        coverage_body.addWidget(QLabel("C=current / S=stale / M=missing / I=invalid. Counts describe configured vendor files, not supported releases."))
        layout.addWidget(coverage_card)
        stats = QHBoxLayout()
        self.feed_stats = {}
        for label, color in (("Current", "#10B981"), ("Stale", "#F59E0B"), ("Missing", "#EF4444"), ("Invalid", "#8B5CF6")):
            card = StatCard(label, color)
            self.feed_stats[label.lower()] = card.value
            stats.addWidget(card)
        layout.addLayout(stats)

        controls, body = self._card("Feed repository", "Freshness uses each feed's validated file timestamp and maximum age policy.")
        row = QHBoxLayout()
        self.start_year = QSpinBox()
        self.start_year.setRange(2002, datetime.now().year)
        self.start_year.setValue(2002)
        self.end_year = QSpinBox()
        self.end_year.setRange(2002, datetime.now().year)
        self.end_year.setValue(datetime.now().year)
        self.include_recent_check = QCheckBox("Recent / modified")
        self.include_recent_check.setChecked(True)
        self.refresh_stale_check = QCheckBox("Refresh stale")
        self.refresh_stale_check.setChecked(True)
        self.remote_check = QCheckBox("ETag / Last-Modified")
        self.remote_check.setChecked(True)
        self.force_check = QCheckBox("Force download")
        self.update_button = QPushButton("Update feeds")
        self.update_button.setProperty("primary", True)
        self.update_button.clicked.connect(self._start_feed_update)
        row.addWidget(QLabel("NVD from"))
        row.addWidget(self.start_year)
        row.addWidget(QLabel("to"))
        row.addWidget(self.end_year)
        row.addWidget(self.include_recent_check)
        row.addStretch(1)
        row.addWidget(self.update_button)
        body.addLayout(row)
        policies = QHBoxLayout()
        policies.addWidget(self.refresh_stale_check)
        policies.addWidget(self.remote_check)
        policies.addWidget(self.force_check)
        policies.addStretch(1)
        body.addLayout(policies)
        self.feed_summary = QLabel("Feed snapshot has not been checked")
        self.feed_summary.setObjectName("inlineStatus")
        self.feed_summary.setWordWrap(True)
        body.addWidget(self.feed_summary)
        self.feed_progress_label = QLabel("Ready to check or update the configured feed snapshot")
        self.feed_progress_label.setObjectName("feedProgressLabel")
        self.feed_progress_label.setWordWrap(True)
        body.addWidget(self.feed_progress_label)
        self.feed_update_progress = QProgressBar()
        self.feed_update_progress.setProperty("prominent", True)
        self.feed_update_progress.setRange(0, 100)
        self.feed_update_progress.setValue(0)
        self.feed_update_progress.setFormat("Ready")
        self.feed_update_progress.setTextVisible(True)
        body.addWidget(self.feed_update_progress)
        layout.addWidget(controls)

        table_card, table_body = self._card("Configured feeds")
        self.feed_table = QTableWidget(0, 5)
        self.feed_table.setHorizontalHeaderLabels(("Feed", "Group", "State", "File updated UTC", "Valid through UTC"))
        self.feed_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.feed_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.feed_table.verticalHeader().setVisible(False)
        header = self.feed_table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        for column in range(1, 5):
            header.setSectionResizeMode(column, QHeaderView.ResizeMode.ResizeToContents)
        table_body.addWidget(self.feed_table)
        layout.addWidget(table_card, 1)
        return page

    def _build_reports_page(self):
        page, layout = self._page()
        card, body = self._card(
            "Audit outputs",
            "Two workbooks per host: (1) OfflineVA CVE evidence  (2) full OS Security Assessment template.",
        )
        self.report_state = QLabel("Complete an audit to populate the report snapshot.")
        self.report_state.setObjectName("monoPanel")
        self.report_state.setWordWrap(True)
        body.addWidget(self.report_state)

        path_grid = QGridLayout()
        path_grid.setHorizontalSpacing(12)
        path_grid.setVerticalSpacing(8)
        self.cve_report_label = QLabel("CVE report: —")
        self.cve_report_label.setObjectName("readinessValue")
        self.cve_report_label.setWordWrap(True)
        self.full_report_label = QLabel("Full assessment: —")
        self.full_report_label.setObjectName("readinessValue")
        self.full_report_label.setWordWrap(True)
        path_grid.addWidget(QLabel("CVE workbook"), 0, 0)
        path_grid.addWidget(self.cve_report_label, 0, 1)
        path_grid.addWidget(QLabel("Full assessment"), 1, 0)
        path_grid.addWidget(self.full_report_label, 1, 1)
        path_grid.setColumnStretch(1, 1)
        body.addLayout(path_grid)

        actions = QHBoxLayout()
        open_cve = QPushButton("Open CVE report")
        open_cve.setProperty("primary", True)
        open_cve.clicked.connect(self._open_cve_report)
        open_full = QPushButton("Open full assessment")
        open_full.setProperty("success", True)
        open_full.clicked.connect(self._open_full_report)
        open_button = QPushButton("Open output folder")
        open_button.setProperty("secondary", True)
        open_button.clicked.connect(self._open_output)
        actions.addWidget(open_cve)
        actions.addWidget(open_full)
        actions.addWidget(open_button)
        actions.addStretch(1)
        body.addLayout(actions)
        layout.addWidget(card)

        guarantees, guarantee_body = self._card("Report guarantees")
        text = QLabel(
            "• Collector TXT (V6 or OfflineVA inventory) is accepted as audit input\n"
            "• CVE workbook: Summary + Vulnerabilities + Inventory (evidence-grade)\n"
            "• Full assessment: enterprise template (Host, Package, Users, Firewall, … + Vulnerabilities)\n"
            "• Confirmed severity totals exclude unresolved product streams\n"
            "• Inventory packages never counted as vulnerabilities without vendor evidence"
        )
        text.setObjectName("readinessValue")
        text.setWordWrap(True)
        guarantee_body.addWidget(text)
        layout.addWidget(guarantees)
        layout.addStretch(1)
        return page


    def _build_settings_page(self):
        page, layout = self._page()
        policy, body = self._card("Detection policy", "The engine remains vendor-first and offline; the UI does not change evidence semantics.")
        text = QLabel(
            "Confirmed Vulnerable requires applicable vendor/local evidence.\n\n"
            "The Target operating system selector constrains the audit to that OS family. Auto-detect uses collector evidence. "
            "Choices marked coverage limited are recorded as an explicit scope, but are not presented as complete vendor-first assurance.\n\n"
            "AUS, E4S, EUS and TUS feed selection uses explicit enabled=true repository configuration; effective runtime state is not collected. "
            "For the base stream, a matching generic RHEL major CPE confirms applicability even when the vendor fixed build is from a newer minor release; unresolved product scope becomes Needs Validation.\n\n"
            "NVD, CISA KEV and EPSS enrich or prioritize an existing vendor decision; they never confirm package exposure by fuzzy CPE matching."
        )
        text.setObjectName("readinessValue")
        text.setWordWrap(True)
        body.addWidget(text)
        layout.addWidget(policy)

        runtime, runtime_body = self._card("Runtime", "The interface follows the Windows app theme automatically.")
        self.runtime_label = QLabel(f"{PRODUCT_NAME} {PRODUCT_VERSION}\nPython {sys.version.split()[0]}\nProject: {self.project_dir}")
        self.runtime_label.setObjectName("monoPanel")
        self.runtime_label.setWordWrap(True)
        runtime_body.addWidget(self.runtime_label)
        layout.addWidget(runtime)
        layout.addStretch(1)
        return page

    # ---------- navigation and styling ----------
    def _show_page(self, index: int):
        titles = (
            ("Overview", "Host, feed and report readiness at a glance"),
            ("Audit", "Collector input, evidence options and live execution"),
            ("Remote", "SSH collection of safe inventory scripts from one or many hosts"),
            ("Feeds", "Validated local sources, freshness and automatic refresh"),
            ("Reports", "Nessus-style advisory findings and inventory evidence"),
            ("Settings", "Applicability policy and runtime information"),
        )
        self.stack.setCurrentIndex(index)
        self.page_title.setText(titles[index][0])
        self.page_subtitle.setText(titles[index][1])
        for button_index, button in enumerate(self.nav_buttons):
            button.setChecked(button_index == index)

    def _apply_theme(self):
        dark = self.dark_theme
        # Material-inspired engineering palette — sharp contrast, clear hierarchy
        if dark:
            c = {
                "bg": "#0B0F19",
                "sidebar": "#0E1422",
                "sidebar_hover": "#161E30",
                "sidebar_active": "#1A2740",
                "card": "#131A2A",
                "surface": "#1A2336",
                "input": "#0C1220",
                "text": "#F0F4FA",
                "muted": "#8B9BB5",
                "border": "#243044",
                "border_soft": "#1C2638",
                "accent": "#3B82F6",
                "accent2": "#2563EB",
                "accent_soft": "#1E3A5F",
                "select": "#1E3A5F",
                "mono": "#0A0E18",
                "success": "#10B981",
                "success_bg": "#0D2A22",
                "success_text": "#34D399",
                "danger": "#EF4444",
                "danger_bg": "#2A1218",
                "danger_text": "#F87171",
                "warn_bg": "#2A2110",
                "warn_text": "#FBBF24",
                "nav_text": "#9AA8C0",
                "nav_active": "#FFFFFF",
                "font": "'Inter', 'Segoe UI Variable', 'Segoe UI', 'SF Pro Display', 'Helvetica Neue', sans-serif",
                "mono_font": "'JetBrains Mono', 'Cascadia Mono', 'SF Mono', 'Consolas', monospace",
            }
        else:
            c = {
                "bg": "#F4F6FA",
                "sidebar": "#0F172A",
                "sidebar_hover": "#1E293B",
                "sidebar_active": "#1E3A5F",
                "card": "#FFFFFF",
                "surface": "#F8FAFC",
                "input": "#FFFFFF",
                "text": "#0F172A",
                "muted": "#64748B",
                "border": "#E2E8F0",
                "border_soft": "#F1F5F9",
                "accent": "#2563EB",
                "accent2": "#1D4ED8",
                "accent_soft": "#DBEAFE",
                "select": "#DBEAFE",
                "mono": "#F8FAFC",
                "success": "#059669",
                "success_bg": "#ECFDF5",
                "success_text": "#047857",
                "danger": "#DC2626",
                "danger_bg": "#FEF2F2",
                "danger_text": "#B91C1C",
                "warn_bg": "#FFFBEB",
                "warn_text": "#B45309",
                "nav_text": "#94A3B8",
                "nav_active": "#FFFFFF",
                "font": "'Inter', 'Segoe UI Variable', 'Segoe UI', 'SF Pro Display', 'Helvetica Neue', sans-serif",
                "mono_font": "'JetBrains Mono', 'Cascadia Mono', 'SF Mono', 'Consolas', monospace",
            }

        self.setStyleSheet(f"""
            /* —— base —— */
            * {{
                font-family: {c['font']};
                font-size: 13px;
                color: {c['text']};
            }}
            QMainWindow, #workspace {{
                background: {c['bg']};
            }}
            QScrollArea, #pageContent {{
                background: {c['bg']};
                border: 0;
            }}

            /* —— sidebar —— */
            #sidebar {{
                background: {c['sidebar']};
                border: 0;
                border-right: 1px solid {'#1A2336' if dark else 'transparent'};
            }}
            #brandMark {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 #3B82F6, stop:1 #1D4ED8);
                color: white;
                border-radius: 11px;
                font-size: 14px;
                font-weight: 700;
                letter-spacing: -0.5px;
            }}
            #brandName {{
                color: white;
                font-size: 15px;
                font-weight: 700;
                letter-spacing: -0.3px;
            }}
            #brandEdition {{
                color: #7B8BA5;
                font-size: 11px;
                font-weight: 500;
            }}
            #navSection {{
                color: #5B6B85;
                font-size: 10px;
                font-weight: 700;
                letter-spacing: 1.2px;
                padding: 0 10px 2px 10px;
            }}
            #sideVersion {{
                color: #5B6B85;
                font-size: 11px;
                font-weight: 500;
            }}
            QPushButton[nav="true"] {{
                text-align: left;
                background: transparent;
                color: {c['nav_text']};
                border: 0;
                border-radius: 10px;
                padding: 10px 14px;
                font-size: 13px;
                font-weight: 550;
            }}
            QPushButton[nav="true"]:hover {{
                background: {c['sidebar_hover']};
                color: {c['nav_active']};
            }}
            QPushButton[nav="true"]:checked {{
                background: {c['sidebar_active']};
                color: {c['nav_active']};
                font-weight: 650;
                border-left: 3px solid {c['accent']};
                padding-left: 11px;
            }}
            #safetyCard {{
                background: {'#0C1F1A' if dark else '#0C1F1A'};
                border: 1px solid {'#1A3D34' if dark else '#1A3D34'};
                border-radius: 12px;
            }}
            #safeTitle {{
                color: #34D399;
                font-size: 11px;
                font-weight: 700;
                letter-spacing: 0.4px;
            }}
            #safeText {{
                color: #7BA89A;
                font-size: 11px;
                line-height: 1.35;
            }}

            /* —— page chrome —— */
            #pageTitle {{
                font-size: 26px;
                font-weight: 700;
                letter-spacing: -0.6px;
                color: {c['text']};
            }}
            #pageSubtitle {{
                color: {c['muted']};
                font-size: 13px;
                font-weight: 450;
            }}
            #footerText {{
                color: {c['muted']};
                font-size: 12px;
            }}

            /* —— cards —— */
            QFrame[card="true"], #heroCard {{
                background: {c['card']};
                border: 1px solid {c['border']};
                border-radius: 14px;
            }}
            QFrame[statCard="true"] {{
                background: {c['card']};
                border: 1px solid {c['border']};
                border-radius: 12px;
                min-width: 110px;
            }}
            #cardTitle {{
                font-size: 15px;
                font-weight: 650;
                letter-spacing: -0.2px;
            }}
            #cardSubtitle, #fieldLabel {{
                color: {c['muted']};
                font-size: 12px;
                font-weight: 450;
            }}
            #statLabel {{
                color: {c['muted']};
                font-size: 10px;
                font-weight: 700;
                letter-spacing: 0.8px;
            }}
            #statValue {{
                font-size: 24px;
                font-weight: 700;
                letter-spacing: -0.5px;
            }}
            #readinessValue {{
                font-size: 13px;
                line-height: 1.5;
            }}
            #inlineStatus {{
                color: {c['muted']};
                font-size: 12px;
            }}

            /* —— status pills —— */
            #warningPill {{
                background: {c['warn_bg']};
                color: {c['warn_text']};
                padding: 8px 14px;
                border-radius: 20px;
                font-size: 12px;
                font-weight: 650;
                border: 1px solid {'#3D3218' if dark else '#FDE68A'};
            }}
            #warningPill[status="success"] {{
                background: {c['success_bg']};
                color: {c['success_text']};
                border: 1px solid {'#1A3D34' if dark else '#A7F3D0'};
            }}
            #warningPill[status="danger"] {{
                background: {c['danger_bg']};
                color: {c['danger_text']};
                border: 1px solid {'#3D1A22' if dark else '#FECACA'};
            }}
            QLabel[status="success"] {{
                background: {c['success_bg']};
                color: {c['success_text']};
                padding: 8px 14px;
                border-radius: 10px;
                font-weight: 600;
                font-size: 12px;
            }}
            QLabel[status="danger"] {{
                background: {c['danger_bg']};
                color: {c['danger_text']};
                padding: 8px 14px;
                border-radius: 10px;
                font-weight: 600;
                font-size: 12px;
            }}

            /* —— form controls —— */
            QLineEdit, QSpinBox, QComboBox, QPlainTextEdit {{
                background: {c['input']};
                border: 1px solid {c['border']};
                border-radius: 10px;
                padding: 9px 12px;
                font-size: 13px;
                selection-background-color: {c['accent']};
                selection-color: white;
            }}
            QLineEdit:hover, QSpinBox:hover, QComboBox:hover, QPlainTextEdit:hover {{
                border: 1px solid {'#3B4A66' if dark else '#CBD5E1'};
            }}
            QLineEdit:focus, QSpinBox:focus, QComboBox:focus, QPlainTextEdit:focus {{
                border: 1.5px solid {c['accent']};
                background: {'#0E1525' if dark else '#FFFFFF'};
            }}
            QComboBox::drop-down {{
                border: 0;
                width: 28px;
            }}
            QComboBox QAbstractItemView {{
                background: {c['card']};
                border: 1px solid {c['border']};
                border-radius: 8px;
                selection-background-color: {c['select']};
                padding: 4px;
            }}
            QCheckBox {{
                spacing: 8px;
                font-size: 13px;
            }}
            QCheckBox::indicator {{
                width: 18px;
                height: 18px;
                border-radius: 5px;
                border: 1.5px solid {c['border']};
                background: {c['input']};
            }}
            QCheckBox::indicator:checked {{
                background: {c['accent']};
                border-color: {c['accent']};
            }}
            QCheckBox::indicator:hover {{
                border-color: {c['accent']};
            }}

            /* —— buttons —— */
            QPushButton {{
                min-height: 36px;
                border-radius: 10px;
                padding: 6px 18px;
                font-size: 13px;
                font-weight: 600;
                border: 0;
            }}
            QPushButton[primary="true"] {{
                background: {c['accent']};
                color: white;
            }}
            QPushButton[primary="true"]:hover {{
                background: {c['accent2']};
            }}
            QPushButton[primary="true"]:pressed {{
                background: {'#1E40AF' if dark else '#1E3A8A'};
            }}
            QPushButton[success="true"] {{
                background: {c['success']};
                color: white;
            }}
            QPushButton[success="true"]:hover {{
                background: {'#059669' if dark else '#047857'};
            }}
            QPushButton[danger="true"] {{
                background: {c['danger_bg']};
                color: {c['danger_text']};
                border: 1px solid {'#4A1F28' if dark else '#FECACA'};
            }}
            QPushButton[danger="true"]:hover {{
                background: {'#3A1520' if dark else '#FEE2E2'};
            }}
            QPushButton[secondary="true"] {{
                background: {c['surface']};
                color: {c['text']};
                border: 1px solid {c['border']};
            }}
            QPushButton[secondary="true"]:hover {{
                background: {'#222C42' if dark else '#F1F5F9'};
                border-color: {'#3B4A66' if dark else '#CBD5E1'};
            }}
            QPushButton[link="true"] {{
                background: transparent;
                color: {c['accent']};
                border: 0;
                padding: 0;
                text-align: left;
                min-height: 0;
            }}
            QPushButton:disabled {{
                color: {c['muted']};
                background: {c['surface']};
                border: 1px solid {c['border_soft']};
            }}

            /* —— tables —— */
            QTableWidget {{
                background: {c['card']};
                alternate-background-color: {c['surface']};
                border: 1px solid {c['border']};
                border-radius: 10px;
                gridline-color: {c['border_soft']};
                selection-background-color: {c['select']};
                selection-color: {c['text']};
                outline: 0;
            }}
            QTableWidget::item {{
                padding: 6px 10px;
            }}
            QHeaderView::section {{
                background: {c['surface']};
                color: {c['muted']};
                border: 0;
                border-bottom: 1px solid {c['border']};
                padding: 10px 12px;
                font-size: 11px;
                font-weight: 700;
                letter-spacing: 0.3px;
                text-transform: uppercase;
            }}

            /* —— progress —— */
            QProgressBar {{
                background: {c['surface']};
                border: 0;
                border-radius: 5px;
                height: 8px;
                text-align: center;
                color: transparent;
            }}
            QProgressBar::chunk {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 {c['accent']}, stop:1 #60A5FA);
                border-radius: 5px;
            }}
            QProgressBar[prominent="true"], QProgressBar[activeFeed="true"] {{
                min-height: 22px;
                max-height: 22px;
                border-radius: 8px;
                color: {c['text']};
                font-weight: 650;
                font-size: 11px;
            }}
            QProgressBar[prominent="true"]::chunk, QProgressBar[activeFeed="true"]::chunk {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 #10B981, stop:1 #34D399);
                border-radius: 8px;
            }}
            #feedProgressLabel {{
                color: {c['text']};
                font-weight: 600;
                font-size: 13px;
            }}

            /* —— mono / code panels —— */
            #monoPanel {{
                background: {c['mono']};
                border: 1px solid {c['border']};
                border-radius: 10px;
                padding: 14px 16px;
                font-family: {c['mono_font']};
                font-size: 12px;
                line-height: 1.45;
                color: {'#C8D4E8' if dark else '#334155'};
            }}

            /* —— toast —— */
            #toast {{
                background: {c['card']};
                border: 1px solid {c['border']};
                border-radius: 14px;
            }}
            #toastTitle {{
                font-weight: 700;
                font-size: 13px;
            }}
            #toastMessage {{
                color: {c['muted']};
                font-size: 12px;
            }}
            #toastDot {{
                color: {c['success']};
                font-size: 14px;
            }}
            #toastClose {{
                background: transparent;
                border: 0;
                color: {c['muted']};
                padding: 0;
                font-size: 16px;
                min-height: 0;
            }}

            /* —— scrollbars —— */
            QScrollBar:vertical {{
                background: transparent;
                width: 8px;
                margin: 4px 2px;
            }}
            QScrollBar::handle:vertical {{
                background: {c['border']};
                border-radius: 4px;
                min-height: 32px;
            }}
            QScrollBar::handle:vertical:hover {{
                background: {'#3B4A66' if dark else '#94A3B8'};
            }}
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                height: 0;
            }}
            QScrollBar:horizontal {{
                background: transparent;
                height: 8px;
                margin: 2px 4px;
            }}
            QScrollBar::handle:horizontal {{
                background: {c['border']};
                border-radius: 4px;
                min-width: 32px;
            }}
            QScrollBar::add-line:horizontal, QScrollBar::sub-line:horizontal {{
                width: 0;
            }}
        """)
        self.theme_label.setText(f"Theme · {'Dark' if dark else 'Light'} (system)")

    def _monitor_system_theme(self):
        current = system_uses_dark_theme()
        if current != self.dark_theme:
            self.dark_theme = current
            self._apply_theme()

    # ---------- path and option state ----------
    def _selected_os_family(self) -> str:
        return str(self.os_family_combo.currentData() or "auto")

    def _on_os_family_changed(self, _index: int = -1):
        family = self._selected_os_family()
        settings = QSettings("OfflineVA", f"{PRODUCT_NAME}-3.6")
        settings.setValue("audit/os_family", family)
        settings.sync()
        supported = {value for value, _label, _tooltip in SUPPORTED_OS_FAMILIES}
        if family == "auto":
            hint = "OS identity will be read once from the collector and used to select the matching vendor evidence."
        elif family in supported:
            hint = f"Only {self.os_family_combo.currentText()} OS evidence will be evaluated for this audit."
        else:
            hint = "This family will be recorded as the requested scope; full vendor-first coverage is not available yet."
        self.os_scope_hint.setText(hint)
        if hasattr(self, "collector_edit"):
            self.collector_edit.setPlainText(collector_command(family))
            self.collector_toggle.setText(collector_label(family))
        self._refresh_coverage_hint()
        if hasattr(self, "feed_table"):
            self._refresh_feed_health()

    def _copy_collector_command(self):
        QApplication.clipboard().setText(self.collector_edit.toPlainText())
        self._toast("Command copied", "دستور کامل کپی شد؛ آن را روی میزبان Linux هدف اجرا کنید.")

    def _toggle_metadata_context(self, enabled: bool):
        for widget in self.context_widgets:
            widget.setVisible(enabled)
        if not enabled:
            self.context_edit.clear()

    def _suggest_metadata_context(self):
        self.coverage_inventory = None
        self._refresh_coverage_hint()
        path = Path(self.input_edit.text().strip())
        if path.is_dir() and not self.is_busy:
            self.context_toggle.setChecked(False)
            return
        if self.is_busy or not path.is_file():
            return
        try:
            from .inventory import parse_inventory
            inventory = parse_inventory(path)
            self.coverage_inventory = inventory
            self._refresh_coverage_hint()
            self.context_toggle.setChecked(any(not pkg.metadata_complete for pkg in inventory.packages))
        except (OSError, ValueError):
            return  # Normal audit validation will show the file error.

    def _browse_input_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Select collector inventory", "", "Collector text (*.txt);;All files (*.*)")
        if path:
            self.input_edit.setText(path)
            self._suggest_metadata_context()
            output = Path(self.output_edit.text())
            if not output.suffix:
                self.output_edit.setText(str(output / f"{Path(path).stem}_cve_report.xlsx"))

    def _browse_input_folder(self):
        path = QFileDialog.getExistingDirectory(self, "Select collector folder")
        if path:
            self.input_edit.setText(path)
            self._suggest_metadata_context()
            output = Path(self.output_edit.text())
            if output.suffix:
                self.output_edit.setText(str(output.parent))

    def _browse_context(self):
        path, _ = QFileDialog.getOpenFileName(self, "Full collector from the same host", "", "Collector text (*.txt)")
        if path:
            self.context_edit.setText(path)

    def _browse_feeds(self):
        path = QFileDialog.getExistingDirectory(self, "Select feed repository")
        if path:
            self.feeds_edit.setText(path)
            self._refresh_feed_health()

    def _browse_output_file(self):
        path, _ = QFileDialog.getSaveFileName(self, "Save audit report", "", "Excel workbook (*.xlsx)")
        if path:
            self.output_edit.setText(path if path.lower().endswith(".xlsx") else path + ".xlsx")

    def _browse_output_folder(self):
        path = QFileDialog.getExistingDirectory(self, "Select output folder")
        if path:
            self.output_edit.setText(path)

    def _feed_options(self) -> dict:
        start, end = self.start_year.value(), self.end_year.value()
        if end < start:
            raise ValueError("NVD end year cannot be earlier than start year.")
        return {
            "root": self.feeds_edit.text().strip(),
            "start": start,
            "end": end,
            "include_recent": self.include_recent_check.isChecked(),
            "force": self.force_check.isChecked(),
            "refresh_stale": self.refresh_stale_check.isChecked(),
            "remote_check": self.remote_check.isChecked(),
            "target_os_family": self._selected_os_family(),
        }

    def _audit_options(self) -> dict:
        input_path = Path(self.input_edit.text().strip())
        if not input_path.exists():
            raise ValueError("Inventory input does not exist.")
        context_path = self.context_edit.text().strip()
        if context_path and (input_path.is_dir() or not Path(context_path).is_file()):
            raise ValueError("Metadata collector must be an existing full collector for a single input host.")
        output = Path(self.output_edit.text().strip())
        if input_path.is_dir() and output.suffix:
            raise ValueError("Folder input requires an output folder, not one XLSX filename.")
        if input_path.is_file() and output.suffix.lower() != ".xlsx":
            output = output / f"{input_path.stem}_cve_report.xlsx"
            self.output_edit.setText(str(output))
        feed = self._feed_options()
        return {
            "input": str(input_path), "feeds": feed["root"], "output": str(output),
            "start": feed["start"], "end": feed["end"],
            "include_recent": feed["include_recent"], "enrich": self.enrich_check.isChecked(),
            "os_family": self._selected_os_family(),
            "inventory_context": context_path,
        }

    def _sync_overview(self):
        input_path = Path(self.input_edit.text().strip())
        self.overview_input.setText(str(input_path) if input_path.exists() else "Select a valid collector TXT or folder")
        if self.last_report or self.last_full_report:
            lines = []
            if self.last_report:
                lines.append(f"CVE: {self.last_report.name}")
            if self.last_full_report:
                lines.append(f"Full: {self.last_full_report.name}")
            self.overview_output.setText("\n".join(lines))
        else:
            self.overview_output.setText(self.output_edit.text().strip() or "Select an output destination")

    # ---------- feed lifecycle ----------
    def _refresh_feed_health(self):
        try:
            options = self._feed_options()
            specs = all_feed_specs(
                options["root"], options["start"], options["end"], options["include_recent"],
                options["target_os_family"],
            )
            self.feed_rows = feed_freshness(options["root"], specs)
            summary = summarize_feed_freshness(self.feed_rows)
        except Exception as exc:
            self.feed_summary.setText(str(exc))
            return
        counts = summary["counts"]
        for state, label in self.feed_stats.items():
            label.setText(str(counts[state]))
        self.feed_table.setRowCount(len(self.feed_rows))
        ordered = sorted(self.feed_rows, key=lambda row: (row.get("state") == "current", row.get("group", ""), row.get("label", "")))
        for row_index, row in enumerate(ordered):
            values = (
                row.get("label", ""), row.get("group", ""), row.get("state", ""),
                row.get("file_updated_at", ""), row.get("valid_until", ""),
            )
            for column, value in enumerate(values):
                self.feed_table.setItem(row_index, column, QTableWidgetItem(str(value)))
        if summary["current"]:
            text = f"●  {counts['current']} configured feeds current"
            detail = f"Configured feeds current (NOT OS coverage) • earliest valid-through {summary['valid_until'] or 'Unknown'}"
            self.feed_pill.setProperty("status", "success")
        else:
            text = f"●  {summary['issue_count']} feeds need attention"
            detail = (
                f"Current {counts['current']} • Stale {counts['stale']} • Missing {counts['missing']} • Invalid {counts['invalid']} • "
                f"earliest validity {summary['valid_until'] or 'Unknown'}"
            )
            self.feed_pill.setProperty("status", "danger")
        self.feed_pill.setText(text)
        self.feed_pill.style().unpolish(self.feed_pill)
        self.feed_pill.style().polish(self.feed_pill)
        self.feed_summary.setText(detail)
        self.overview_feed.setText(detail)
        self._refresh_coverage_hint()
        rows = coverage_rows(options["root"])
        self.coverage_table.setRowCount(len(rows))
        for index, row in enumerate(rows):
            counts = row["feed_state"]
            state = "/".join(str(counts.get(key, 0)) for key in ("current", "stale", "missing", "invalid")) if counts else "Not configured"
            for col, value in enumerate((row["distribution"], row["versions"], row["architecture"], row["implementation"], state, row["note"])):
                item = QTableWidgetItem(value)
                item.setToolTip(row["architecture"] + "\n" + row["note"])
                self.coverage_table.setItem(index, col, item)
        self.coverage_table.resizeRowsToContents()

    def _refresh_coverage_hint(self):
        if hasattr(self, "coverage_hint"):
            root = self.feeds_edit.text() if hasattr(self, "feeds_edit") else self.project_dir / "feeds"
            self.coverage_hint.setText(scope_preview(self._selected_os_family(), root, self.coverage_inventory))

    def _start_feed_update(self, _checked=False, audit_options: dict | None = None):
        if self.is_busy:
            return
        try:
            options = self._feed_options()
        except ValueError as exc:
            QMessageBox.warning(self, "Feed options", str(exc))
            return
        specs = all_feed_specs(
            options["root"], options["start"], options["end"], options["include_recent"],
            options["target_os_family"],
        )
        total_feeds = max(1, len(specs))
        self.cancel_event.clear()
        self._set_busy(True, "Updating feeds")
        self.signals.feed_progress.emit(0.0, "Preparing feed update", f"0 of {len(specs)} feeds")

        def worker():
            completed = 0

            def update_label(message: str):
                nonlocal completed
                self.signals.log.emit(message)
                payload = message.partition("] ")[2] or message
                if message.startswith("[CHECK]"):
                    self.signals.feed_progress.emit(
                        completed * 100.0 / total_feeds,
                        payload,
                        f"Checking feed {min(completed + 1, total_feeds)} of {total_feeds}",
                    )
                elif message.startswith("[+] Downloading"):
                    label = payload.removeprefix("Downloading ").split(" (attempt", 1)[0]
                    self.signals.feed_progress.emit(
                        completed * 100.0 / total_feeds,
                        label,
                        f"Downloading feed {min(completed + 1, total_feeds)} of {total_feeds}",
                    )
                elif message.startswith("[OK]") or (message.startswith("[WARN]") and "retrying" not in message.lower()):
                    completed = min(completed + 1, total_feeds)
                    label = payload.split(":", 1)[0]
                    self.signals.feed_progress.emit(
                        completed * 100.0 / total_feeds,
                        label,
                        f"Completed {completed} of {total_feeds} feeds",
                    )

            def download_progress(label: str, downloaded: int, total: int):
                file_fraction = min(1.0, downloaded / total) if total else 0.0
                overall = (completed + file_fraction) * 100.0 / total_feeds
                downloaded_mb = downloaded / (1024 * 1024)
                if total:
                    detail = f"Feed {min(completed + 1, total_feeds)} of {total_feeds} • {downloaded_mb:.1f} / {total / (1024 * 1024):.1f} MB"
                else:
                    detail = f"Feed {min(completed + 1, total_feeds)} of {total_feeds} • {downloaded_mb:.1f} MB downloaded"
                self.signals.feed_progress.emit(overall, label, detail)

            try:
                statuses = bootstrap_all(
                    options["root"], options["start"], options["end"], options["include_recent"],
                    target_os_family=options["target_os_family"],
                    force=options["force"], refresh_stale=options["refresh_stale"],
                    remote_check=options["remote_check"], cancel=self.cancel_event.is_set,
                    log=update_label, progress=download_progress,
                )
                self.signals.feed_done.emit(statuses, audit_options)
            except OperationCancelled:
                self.signals.stopped.emit("Feed update cancelled")
            except Exception as exc:
                self.signals.operation_error.emit("Feed update failed", str(exc))

        threading.Thread(target=worker, daemon=True).start()

    def _on_feed_progress(self, percent: float, label: str, detail: str):
        value = max(0, min(100, int(round(percent))))
        self.feed_progress_label.setText(f"{label}  —  {detail}")
        self.feed_update_progress.setRange(0, 100)
        self.feed_update_progress.setValue(value)
        self.feed_update_progress.setFormat(f"{value}%")
        self.progress.setRange(0, 100)
        self.progress.setValue(value)
        self.progress.setProperty("activeFeed", True)
        self.progress.style().unpolish(self.progress)
        self.progress.style().polish(self.progress)
        self.progress.setTextVisible(True)
        self.progress.setFormat(f"{value}%  •  {label}")
        self.audit_ready.setText(f"Updating feeds — {detail}")
        self.status_label.setText(f"Updating feeds — {value}%")

    def _on_feed_done(self, statuses, audit_options):
        failures = [row for row in statuses if str(getattr(row, "state", "")).startswith("error") or getattr(row, "state", "") == "invalid"]
        if failures:
            message = f"{len(failures)} feed(s) failed. The existing snapshot may be incomplete or stale."
            self._refresh_feed_health()
            self._set_busy(False, "Feed update incomplete")
            self.feed_progress_label.setText(message)
            self.feed_update_progress.setFormat("Incomplete")
            self._toast("Feed update incomplete", message, "warning")
            if audit_options is not None:
                answer = QMessageBox.question(self, "Feed update incomplete", message + "\n\nContinue the audit using the available feeds and record the limitations?", QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No)
                if answer == QMessageBox.StandardButton.Yes:
                    self._start_audit_process(audit_options)
            return
        self._on_feed_progress(100.0, "Feed snapshot ready", f"Completed {len(statuses)} configured feeds")
        self._refresh_feed_health()
        self._toast("Feed snapshot ready", f"{len(statuses)} configured feed records checked.")
        if audit_options is not None:
            self._start_audit_process(audit_options)
        else:
            self._set_busy(False, "Ready")

    # ---------- audit lifecycle ----------
    def _request_audit(self):
        if self.is_busy:
            return
        try:
            options = self._audit_options()
            family = options.get("os_family") or "auto"
            limited = {key for key, _ in LIMITED_OS_FAMILIES}
            if family in limited:
                answer = QMessageBox.warning(
                    self,
                    "Assessment scope limited",
                    (
                        f"OS family {family!r} has no full vendor adapter.\n\n"
                        "The run will produce inventory and coverage notes only; "
                        "Confirmed Vulnerable findings from vendor evidence will not be available.\n\n"
                        "Continue anyway?"
                    ),
                    QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                    QMessageBox.StandardButton.No,
                )
                if answer != QMessageBox.StandardButton.Yes:
                    return
            if Path(options["input"]).is_file() and not options["inventory_context"]:
                from .inventory import parse_inventory

                inventory = parse_inventory(options["input"])
                if not inventory.packages or any(not pkg.metadata_complete for pkg in inventory.packages):
                    self._start_audit_process(options)
                    self._append_log("[WARN] Inventory-only report: select the matching full Metadata collector to enable vulnerability assessment. Feed refresh cannot supply missing host metadata.")
                    return
            feed_options = self._feed_options()
            specs = all_feed_specs(
                feed_options["root"], feed_options["start"], feed_options["end"], feed_options["include_recent"],
                options["os_family"],
            )
            freshness = summarize_feed_freshness(feed_freshness(feed_options["root"], specs))
        except ValueError as exc:
            QMessageBox.warning(self, "Audit configuration", str(exc))
            return

        refresh = self.auto_feeds_check.isChecked()
        if not freshness["current"] and not refresh:
            counts = freshness["counts"]
            answer = QMessageBox.question(
                self,
                "Feed snapshot needs attention",
                f"Stale: {counts['stale']}\nMissing: {counts['missing']}\nInvalid: {counts['invalid']}\n"
                f"Earliest validity: {freshness['valid_until'] or 'Unknown'}\n\n"
                "Update automatically before this audit?\n\nChoose No to continue and record the existing snapshot in the report.",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.Yes,
            )
            refresh = answer == QMessageBox.StandardButton.Yes
        if refresh:
            self._start_feed_update(audit_options=options)
        else:
            self._start_audit_process(options)

    def _start_audit_process(self, options: dict):
        self.cancel_event.clear()
        self._set_busy(True, "Running vendor evidence audit")
        self.log_edit.clear()
        self._append_log(
            f"[SCOPE] Target OS family: {options.get('os_family', 'auto')} "
            f"({self.os_family_combo.currentText()})"
        )
        arguments = self._audit_arguments(options)
        self.process = QProcess(self)
        self.process.setProcessChannelMode(QProcess.ProcessChannelMode.MergedChannels)
        self.process.setWorkingDirectory(str(self.project_dir))
        self.process.readyReadStandardOutput.connect(self._read_process_output)
        self.process.errorOccurred.connect(self._audit_process_error)
        self.process.finished.connect(lambda code, _status: self._audit_finished(code, options))
        self.process.start(sys.executable, arguments)

    def _audit_arguments(self, options: dict) -> list[str]:
        arguments = [
            "-u", "-X", "utf8", str(self.project_dir / "cli.py"), "audit",
            "--input", options["input"], "--feeds", options["feeds"], "--output", options["output"],
            "--start-year", str(options["start"]), "--end-year", str(options["end"]),
            "--feed-policy", "use-existing",
            "--os-family", options.get("os_family", "auto"),
        ]
        if not options["include_recent"]:
            arguments.append("--no-recent")
        if not options["enrich"]:
            arguments.append("--no-enrich")
        if options.get("inventory_context"):
            arguments.extend(["--inventory-context", options["inventory_context"]])
        return arguments

    def _audit_process_error(self, error):
        if error == QProcess.ProcessError.FailedToStart:
            detail = self.process.errorString() if self.process else "Could not start the audit process"
            self.process = None
            self._on_operation_error("Audit failed to start", detail)

    def _read_process_output(self):
        if self.process:
            text = bytes(self.process.readAllStandardOutput()).decode("utf-8", errors="replace")
            for line in text.splitlines():
                self._append_log(line)

    def _audit_finished(self, code: int, options: dict):
        self._read_process_output()
        self.process = None
        if self.cancel_event.is_set():
            self._on_stopped("Audit cancelled")
            return
        if code not in {0, 3}:
            self._on_operation_error("Audit failed", f"Audit process exited with code {code}")
            return
        output = Path(options["output"])
        self.last_report = None
        self.last_full_report = None
        if output.is_dir():
            cve_reports = sorted(
                [p for p in output.glob("*_cve_report.xlsx")],
                key=lambda path: path.stat().st_mtime,
                reverse=True,
            )
            full_reports = sorted(
                [p for p in output.glob("*_OS_Security_Assessment.xlsx")],
                key=lambda path: path.stat().st_mtime,
                reverse=True,
            )
            self.last_report = cve_reports[0] if cve_reports else None
            self.last_full_report = full_reports[0] if full_reports else None
            if self.last_report is None:
                other = sorted(
                    [p for p in output.glob("*.xlsx") if "OS_Security_Assessment" not in p.name],
                    key=lambda path: path.stat().st_mtime,
                    reverse=True,
                )
                self.last_report = other[0] if other else None
        else:
            self.last_report = output
            candidate = output.with_name(
                output.stem.replace("_cve_report", "") + "_OS_Security_Assessment.xlsx"
            )
            if candidate.is_file():
                self.last_full_report = candidate
            else:
                alt = output.with_name(output.stem + "_OS_Security_Assessment.xlsx")
                self.last_full_report = alt if alt.is_file() else None
        values = self._load_report_snapshot() or {}
        status = str(values.get("Analysis Status", "Report could not be verified"))
        inferred = "partial" if status.startswith("Partial Assessment") else "complete" if status.startswith("Completed") else "incomplete"
        assessment = values.get("Assessment State", inferred)
        limited = code == 3 or assessment != "complete" or status != "Completed"
        title = "Assessment incomplete" if assessment == "incomplete" else "Partial assessment" if limited else "Audit complete"
        self._set_busy(False, title)
        self._sync_overview()
        detail = str(values.get("Assessment Limitation") or values.get("Finding Validity") or status)
        cve_name = Path(str(self.last_report)).name if self.last_report else "—"
        full_name = Path(str(self.last_full_report)).name if self.last_full_report else "—"
        self.report_state.setText(
            f"{title}\n\nCVE report: {self.last_report or '—'}\n"
            f"Full assessment: {self.last_full_report or '—'}\n\n{status}\n{detail}"
        )
        if hasattr(self, "cve_report_label"):
            self.cve_report_label.setText(str(self.last_report) if self.last_report else "—")
        if hasattr(self, "full_report_label"):
            self.full_report_label.setText(str(self.last_full_report) if self.last_full_report else "—")
        toast_detail = f"{cve_name}"
        if self.last_full_report:
            toast_detail += f"\n+ {full_name}"
        self._toast(title, f"{toast_detail}\n{detail}", "warning" if limited else "success")
        if self.last_full_report:
            self._append_log(f"[OK] Full OS Security Assessment: {self.last_full_report}")
        if self.last_report:
            self._append_log(f"[OK] CVE evidence report: {self.last_report}")

    def _load_report_snapshot(self):
        if not self.last_report or not self.last_report.exists():
            return
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(self.last_report, read_only=True, data_only=True)
            values = {row[0]: row[1] for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)}
            workbook.close()
            mapping = {
                "Confirmed": "Vulnerability Findings", "Critical": "Critical", "High": "High",
                "Medium": "Medium", "Low": "Low", "Review": "Scope / Evidence Review Candidates",
            }
            for label, key in mapping.items():
                self.report_stats[label].setText(str(values.get(key, 0)))
            return values
        except Exception as exc:
            self._append_log(f"[WARN] Could not read report summary: {exc}")

    def _stop_operation(self):
        self.cancel_event.set()
        if self.process and self.process.state() != QProcess.ProcessState.NotRunning:
            self.process.kill()
        self.status_label.setText("Cancellation requested")

    def _set_busy(self, busy: bool, text: str):
        self.is_busy = busy
        self.run_button.setEnabled(not busy)
        self.update_button.setEnabled(not busy)
        self.os_family_combo.setEnabled(not busy)
        self.stop_button.setEnabled(busy)
        self.audit_ready.setText(text)
        self.status_label.setText(text)
        self.progress.setProperty("activeFeed", False)
        self.progress.style().unpolish(self.progress)
        self.progress.style().polish(self.progress)
        if busy:
            self.progress.setRange(0, 0)
            self.progress.setTextVisible(False)
        else:
            self.progress.setRange(0, 1)
            self.progress.setValue(1)
            self.progress.setTextVisible(False)

    def _append_log(self, text: str):
        stamp = datetime.now().strftime("%H:%M:%S")
        self.log_edit.appendPlainText(f"[{stamp}] {text}")

    def _on_operation_error(self, title: str, message: str):
        self._set_busy(False, message)
        if "feed" in title.lower():
            self.feed_progress_label.setText(message)
            self.feed_update_progress.setFormat("Failed")
        QMessageBox.critical(self, title, message)

    def _on_stopped(self, message: str):
        self._set_busy(False, message)
        self.feed_progress_label.setText(message)
        self.feed_update_progress.setFormat("Stopped")
        self._toast("Operation stopped", message, "warning")

    def _open_path_file(self, path: Path | None, label: str) -> None:
        if not path or not Path(path).exists():
            QMessageBox.information(self, label, f"No {label.lower()} available yet. Run an audit first.")
            return
        path = Path(path)
        try:
            os.startfile(path)  # type: ignore[attr-defined]
        except Exception as exc:
            QMessageBox.information(self, label, f"{path}\n\n{exc}")

    def _open_cve_report(self):
        self._open_path_file(self.last_report, "CVE report")

    def _open_full_report(self):
        self._open_path_file(self.last_full_report, "Full assessment")

    def _open_output(self):
        path = self.last_report or Path(self.output_edit.text().strip())
        folder = path if path.is_dir() else path.parent
        folder.mkdir(parents=True, exist_ok=True)
        try:
            os.startfile(folder)  # type: ignore[attr-defined]
        except Exception as exc:
            QMessageBox.information(self, "Output folder", f"{folder}\n\n{exc}")

    # ---------- remote SSH collection ----------
    def _remote_browse_key(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select private key", "", "Private key (*);;PEM (*.pem);;All files (*.*)"
        )
        if path:
            self.remote_key_path.setText(path)

    def _remote_browse_output(self):
        path = QFileDialog.getExistingDirectory(self, "Inventory output folder")
        if path:
            self.remote_output_dir.setText(path)

    def _remote_selected_family(self) -> str:
        return str(self.remote_default_family.currentData() or "auto")

    def _remote_append_log(self, message: str):
        self.remote_log.appendPlainText(message)
        self.remote_log.verticalScrollBar().setValue(self.remote_log.verticalScrollBar().maximum())

    def _remote_export_bash(self):
        family = self._remote_selected_family()
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Export collector script",
            f"collector_{family}.sh",
            "Shell script (*.sh);;All files (*.*)",
        )
        if not path:
            return
        try:
            export_collector_bash(family, Path(path))
            self._toast("Collector exported", str(path))
            self._remote_append_log(f"Exported collector for family={family} → {path}")
        except Exception as exc:
            QMessageBox.warning(self, "Export failed", str(exc))

    def _remote_copy_bash(self):
        family = self._remote_selected_family()
        try:
            text = collector_command(family)
            QApplication.clipboard().setText(text)
            self._toast("Copied", f"Collector bash for {family} is on the clipboard")
        except Exception as exc:
            QMessageBox.warning(self, "Copy failed", str(exc))

    def _remote_build_targets(self) -> list[Target]:
        text = self.remote_targets_edit.toPlainText()
        targets = parse_target_list(
            text,
            default_user=self.remote_default_user.text().strip() or "root",
            default_port=int(self.remote_default_port.value()),
            default_family=self._remote_selected_family(),
        )
        method = self.remote_auth_method.currentData()
        password = self.remote_password.text() or None
        key_path = self.remote_key_path.text().strip() or None
        key_pass = self.remote_key_pass.text() or None
        for t in targets:
            if method == "password":
                t.password = password
                t.private_key_path = None
            elif method == "key":
                t.private_key_path = key_path
                t.private_key_passphrase = None
                t.password = password  # optional fallback
            else:  # key_pass
                t.private_key_path = key_path
                t.private_key_passphrase = key_pass
                t.password = password
        return targets

    def _remote_build_jump(self) -> JumpHost | None:
        if not self.remote_use_jump.isChecked():
            return None
        host = self.remote_jump_host.text().strip()
        if not host:
            return None
        return JumpHost(
            host=host,
            port=int(self.remote_jump_port.value()),
            username=self.remote_jump_user.text().strip() or "root",
            password=self.remote_jump_password.text() or None,
            private_key_path=self.remote_jump_key.text().strip() or None,
        )

    def _remote_build_opts(self) -> SSHOptions:
        policy = "auto_add" if self.remote_auto_hostkey.isChecked() else "warn"
        return SSHOptions(
            timeout=30.0,
            command_timeout=float(self.remote_cmd_timeout.value()),
            host_key_policy=policy,
            allow_legacy=self.remote_legacy.isChecked(),
            prefer_paramiko=True,
        )

    def _remote_stop(self):
        self._remote_stop_event.set()
        self.remote_status.setText("Stopping…")
        self._remote_append_log("Stop requested — waiting for in-flight host to finish")

    def _remote_start(self):
        if self._remote_thread and self._remote_thread.is_alive():
            QMessageBox.information(self, "Busy", "A remote collection is already running.")
            return
        try:
            targets = self._remote_build_targets()
        except Exception as exc:
            QMessageBox.warning(self, "Targets", str(exc))
            return
        if not targets:
            QMessageBox.warning(self, "Targets", "Add at least one host (one per line).")
            return
        out_dir = Path(self.remote_output_dir.text().strip() or ".")
        try:
            out_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            QMessageBox.warning(self, "Output folder", str(exc))
            return

        opts = self._remote_build_opts()
        jump = self._remote_build_jump()
        workers = int(self.remote_workers.value())
        self._remote_stop_event.clear()
        self.remote_run_btn.setEnabled(False)
        self.remote_stop_btn.setEnabled(True)
        self.remote_progress.setValue(0)
        self.remote_status.setText(f"Starting {len(targets)} host(s)…")
        self._remote_append_log(
            f"=== SSH collect start  hosts={len(targets)}  workers={workers}  "
            f"legacy={opts.allow_legacy}  jump={jump.host if jump else 'none'} ==="
        )

        def worker():
            def log_fn(msg: str):
                QTimer.singleShot(0, lambda m=msg: self._remote_append_log(m))

            def progress_fn(host: str, status: str, frac: float):
                def _ui():
                    self.remote_progress.setValue(int(frac * 100))
                    self.remote_status.setText(f"{host}: {status}")
                QTimer.singleShot(0, _ui)

            results = collect_many(
                targets,
                out_dir,
                opts=opts,
                jump=jump,
                log=log_fn,
                progress=progress_fn,
                stop_event=self._remote_stop_event,
                max_workers=workers,
            )

            def done():
                ok = sum(1 for r in results if r.success)
                fail = len(results) - ok
                self.remote_progress.setValue(100)
                self.remote_run_btn.setEnabled(True)
                self.remote_stop_btn.setEnabled(False)
                self.remote_status.setText(f"Done — {ok} ok, {fail} failed")
                self._remote_append_log(f"=== finished  ok={ok}  fail={fail}  dir={out_dir} ===")
                for r in results:
                    if r.success and r.output_path:
                        self._remote_append_log(f"  OK  {r.target.host} → {r.output_path.name} ({r.bytes_written} B)")
                    else:
                        self._remote_append_log(f"  FAIL {r.target.host}: {r.error}")
                if ok:
                    self._toast("Remote collection finished", f"{ok} inventory file(s) in {out_dir}")
                elif fail:
                    self._toast("Remote collection failed", f"{fail} host(s) failed — see log", "warning")

            QTimer.singleShot(0, done)

        self._remote_thread = threading.Thread(target=worker, daemon=True)
        self._remote_thread.start()

    # ---------- toast ----------
    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._position_toast()

    def _position_toast(self):
        if self.active_toast and self.active_toast.isVisible():
            self.active_toast.adjustSize()
            self.active_toast.move(self.width() - self.active_toast.width() - 28, 24)

    def _toast(self, title: str, message: str, kind: str = "success", duration: int = 5200):
        if self.active_toast:
            self.active_toast.close()
        self.active_toast = Toast(self, title, message, kind)
        self.active_toast.show()
        self.active_toast.raise_()
        self._position_toast()
        QTimer.singleShot(duration, self.active_toast.close)



def run_qt_app() -> int:
    app = QApplication.instance() or QApplication([])
    app.setApplicationName(PRODUCT_NAME)
    app.setOrganizationName("OfflineVA")
    window = OfflineVAWindow()
    window.show()
    return app.exec()



if __name__ == "__main__":
    # python offline_va/gui_qt.py  →  load as package then run
    import sys
    from pathlib import Path
    _root = Path(__file__).resolve().parent.parent
    if str(_root) not in sys.path:
        sys.path.insert(0, str(_root))
    from offline_va.gui_qt import run_qt_app
    raise SystemExit(run_qt_app())
