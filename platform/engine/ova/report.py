from __future__ import annotations

from dataclasses import asdict, is_dataclass
from datetime import datetime
from pathlib import Path
import math
import os
import tempfile

from openpyxl import Workbook
from openpyxl.chart.chartspace import ChartSpace
from openpyxl.chart.reader import read_chart
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.xml.functions import fromstring
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from .engine import AnalysisResult
from .models import FindingStatus

# Client-facing report: intentionally limited to three sheets.
# 1) Summary         : host/result/coverage overview
# 2) Vulnerabilities : actionable findings only
# 3) Inventory       : native package-manager + standalone software inventory

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9E1F2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SEVERITY = {
    "Critical": ("C2212E", "FFFFFF"),
    "High": ("F37314", "FFFFFF"),
    "Medium": ("FAD7A0", "000000"),
    "Low": ("FFFF00", "000000"),
    "Info": ("5B9BD5", "FFFFFF"),
}
STATUS_FILL = {
    FindingStatus.CONFIRMED_VULNERABLE.value: "F4CCCC",
    FindingStatus.VENDOR_AFFECTED_NO_FIX.value: "FCE5CD",
    FindingStatus.POTENTIAL.value: "FFF2CC",
    FindingStatus.NEEDS_VALIDATION.value: "D9D2E9",
    FindingStatus.REBOOT_REQUIRED.value: "FCE5CD",
    FindingStatus.ANALYSIS_INCOMPLETE.value: "E7E6E6",
}

ACTIONABLE_STATUSES = {
    FindingStatus.CONFIRMED_VULNERABLE,
    FindingStatus.VENDOR_AFFECTED_NO_FIX,
    FindingStatus.POTENTIAL,
    FindingStatus.NEEDS_VALIDATION,
    FindingStatus.REBOOT_REQUIRED,
    FindingStatus.ANALYSIS_INCOMPLETE,
}

CONFIRMED_STATUSES = {
    FindingStatus.CONFIRMED_VULNERABLE,
    FindingStatus.VENDOR_AFFECTED_NO_FIX,
    FindingStatus.REBOOT_REQUIRED,
}


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def _finish_sheet(ws, widths: dict[str, float], freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # untrusted inventory/feed text is never an Excel formula
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        lines = max(
            sum(max(1, math.ceil(len(line) / max(10, widths.get(cell.column_letter, 15) - 2))) for line in str(cell.value or "").split("\n"))
            for cell in row
        )
        ws.row_dimensions[row[0].row].height = min(409, max(30, lines * 15 + 8))


def _add_table(ws, name: str) -> None:
    # A structured table owns the filter. Do not add worksheet-level auto_filter.
    if ws.max_row < 2 or ws.max_column < 1:
        return
    table = Table(displayName=name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
    )
    ws.add_table(table)


def _family_tokens(result: AnalysisResult) -> list[str]:
    return {
        "suse": ["suse"],
        "redhat": ["red hat"],
        "redhat-compatible": ["centos", "rocky", "alma", "red hat"],
        "huawei": ["huawei", "euler", "openeuler"],
        "debian": ["debian"],
        "ubuntu": ["ubuntu", "canonical"],
        "oracle-linux": ["oracle"],
        "amazon-linux": ["amazon"],
        "oracle": ["oracle"],
    }.get(result.inventory.os.family, [])


def _third_party_state(result: AnalysisResult, pkg) -> bool | None:
    manager = str(getattr(pkg, "package_manager", "rpm") or "rpm").lower()
    if manager in {"deb", "dpkg"}:
        origin_text = f"{pkg.vendor} {pkg.packager} {pkg.maintainer} {pkg.origin}".lower()
        custom_tokens = ("opennetworklinux", "onl", "sinos", "sina", "accton")
        if any(token in f"{pkg.name} {origin_text}" for token in custom_tokens):
            return True
        if "debian" in origin_text or "@debian.org" in origin_text:
            return False
        return None
    tokens = _family_tokens(result)
    if not tokens:
        return None
    vendor_text = f"{pkg.vendor} {pkg.packager}".lower()
    if not vendor_text.strip():
        return None
    return not any(token in vendor_text for token in tokens)


def _is_third_party_package(result: AnalysisResult, pkg) -> bool:
    return _third_party_state(result, pkg) is True


def _actionable_findings(result: AnalysisResult):
    return [x for x in result.findings if x.status in ACTIONABLE_STATUSES]


def _finding_groups(result: AnalysisResult) -> list[list]:
    """Present one host/advisory plugin result with all packages and CVEs."""
    groups: dict[str, list] = {}
    for item in _actionable_findings(result):
        advisory_key = item.evidence.advisory or item.cve or item.evidence.definition_id
        groups.setdefault(advisory_key, []).append(item)
    severity_order = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "Info": 4}
    return sorted(
        groups.values(),
        key=lambda items: (
            1 if _group_requires_review(items) else 0,
            severity_order.get(_group_severity(items), 5),
            items[0].evidence.advisory or items[0].cve,
        ),
    )


def _group_severity(items: list) -> str:
    order = {"Info": 0, "Low": 1, "Medium": 2, "High": 3, "Critical": 4}
    return max((x.severity if x.severity in order else "Info" for x in items), key=lambda value: order[value], default="Info")


def _group_vendor_severity(items: list) -> str:
    order = {"Info": 0, "Low": 1, "Medium": 2, "High": 3, "Critical": 4}
    values = []
    if all(str((item.evidence.raw or {}).get("vendor_severity")) == "Not rated" for item in items):
        return "Not rated"
    for item in items:
        value = str((item.evidence.raw or {}).get("vendor_severity") or item.severity or "Info")
        values.append(value if value in order else "Info")
    return max(values, key=lambda value: order[value], default="Info")


def _group_requires_review(items: list) -> bool:
    return any(
        item.review_required
        or item.status in {FindingStatus.NEEDS_VALIDATION, FindingStatus.POTENTIAL, FindingStatus.ANALYSIS_INCOMPLETE}
        for item in items
    )


def _group_is_confirmed(items: list) -> bool:
    return not _group_requires_review(items) and any(item.status in CONFIRMED_STATUSES for item in items)


def _group_status(items: list) -> str:
    if any(item.status == FindingStatus.ANALYSIS_INCOMPLETE for item in items):
        return FindingStatus.ANALYSIS_INCOMPLETE.value
    if _group_requires_review(items):
        return FindingStatus.NEEDS_VALIDATION.value
    for status in (
        FindingStatus.CONFIRMED_VULNERABLE,
        FindingStatus.VENDOR_AFFECTED_NO_FIX,
        FindingStatus.REBOOT_REQUIRED,
    ):
        if any(item.status == status for item in items):
            return status.value
    return items[0].status.value


def _group_class(items: list) -> str:
    return "Confirmed Vulnerability" if _group_is_confirmed(items) else "Scope / Evidence Review"


def _coverage_issues(result: AnalysisResult) -> list[str]:
    """Return a concise list of material coverage limitations for the client report."""
    if not result.coverage:
        return ["Coverage metadata was not produced"]
    acceptable = {
        "feed present", "covered", "valid", "available", "complete", "conservative",
        "enrichment only", "prioritization only", "auto-detected", "manually selected",
        "collected", "release identified", "index present", "configured", "compatible",
    }
    issues: list[str] = []
    for item in result.coverage:
        area = str(item.get("Area", "")).strip() or "Coverage"
        status = str(item.get("Status", "")).strip()
        if status.lower() not in acceptable:
            issues.append(f"{area}: {status or 'Unknown'}")
    return issues


def _coverage_text(result: AnalysisResult) -> str:
    issues = _coverage_issues(result)
    return "; ".join(issues) if issues else "No material coverage limitations reported."


def _feed_status_text(feed_status: list | None) -> str:
    if not feed_status:
        return "Not provided"
    current = stale = invalid = missing = other = 0
    for item in feed_status:
        if is_dataclass(item):
            value = asdict(item)
        elif isinstance(item, dict):
            value = item
        else:
            continue
        state = str(value.get("state", "")).lower()
        if state in {"valid", "current", "present", "downloaded"}:
            current += 1
        elif state == "stale":
            stale += 1
        elif state == "invalid":
            invalid += 1
        elif state == "missing":
            missing += 1
        else:
            other += 1
    parts = [f"{current} current"]
    if stale:
        parts.append(f"{stale} stale")
    if invalid:
        parts.append(f"{invalid} invalid")
    if missing:
        parts.append(f"{missing} missing")
    if other:
        parts.append(f"{other} other")
    return ", ".join(parts)


def _relevant_feed_status(result: AnalysisResult, feed_status: list | None) -> list:
    """Exclude unrelated vendor repositories from this host's validity claim."""
    vendor_group = {
        "redhat": "redhat",
        "suse": "suse",
        "debian": "debian",
        "huawei": "huawei",
    }.get(result.inventory.os.family)
    relevant: list = []
    for item in feed_status or []:
        value = asdict(item) if is_dataclass(item) else item if isinstance(item, dict) else {}
        group = str(value.get("group") or "").lower()
        if not group or group in {"nvd", "enrichment"} or (vendor_group and group == vendor_group):
            relevant.append(item)
    return relevant


def _feed_snapshot(feed_status: list | None) -> dict:
    rows: list[dict] = []
    for item in feed_status or []:
        if is_dataclass(item):
            rows.append(asdict(item))
        elif isinstance(item, dict):
            rows.append(item)
    checked = sorted(str(x.get("checked_at") or x.get("updated_at") or "") for x in rows if x.get("checked_at") or x.get("updated_at"))
    expiries = sorted(str(x.get("valid_until") or "") for x in rows if x.get("valid_until"))
    issue_rows = [x for x in rows if str(x.get("state") or "").lower() in {"stale", "invalid", "missing", "error-required", "error-optional"}]
    checked_text = "Not provided"
    if checked:
        checked_text = checked[0] if checked[0] == checked[-1] else f"{checked[0]} — {checked[-1]}"
    return {
        "checked": checked_text,
        "valid_until": expiries[0] if expiries else "Unknown",
        "issues": issue_rows,
        "status": _feed_status_text(feed_status),
    }


def _analysis_status(result: AnalysisResult, findings, feed_status: list | None = None) -> str:
    if result.assessment_state == "incomplete":
        return "Analysis Incomplete - Not Fully Assessed"
    if result.assessment_state == "partial":
        return "Partial Assessment - Coverage Limited"
    if not result.inventory.packages and not result.inventory.software:
        return "Analysis Incomplete - No Usable Inventory"
    if any(x.status == FindingStatus.ANALYSIS_INCOMPLETE for x in findings):
        return "Analysis Incomplete - Review Required"
    if any(x.review_required or x.status in {FindingStatus.NEEDS_VALIDATION, FindingStatus.POTENTIAL} for x in findings):
        return "Completed with Review Items"
    if _feed_snapshot(feed_status)["issues"]:
        return "Completed with Stale/Incomplete Feed Snapshot"
    if _coverage_issues(result):
        return "Completed with Limitations"
    return "Completed"


def _summary(ws, result: AnalysisResult, feed_status: list | None) -> None:
    inv = result.inventory
    findings = _actionable_findings(result)
    finding_groups = _finding_groups(result)
    confirmed_groups = [items for items in finding_groups if _group_is_confirmed(items)]
    review_groups = [items for items in finding_groups if not _group_is_confirmed(items)]
    unique_cves = {x.cve for items in confirmed_groups for x in items if x.cve}
    third_party_count = sum(1 for pkg in inv.packages if _is_third_party_package(result, pkg))
    unknown_origin_count = sum(1 for pkg in inv.packages if _third_party_state(result, pkg) is None)
    managers = sorted({str(getattr(pkg, "package_manager", "rpm") or "rpm").upper() for pkg in inv.packages})
    severity_counts = {key: 0 for key in ("Critical", "High", "Medium", "Low", "Info")}
    for items in confirmed_groups:
        severity_counts[_group_severity(items)] += 1
    relevant_feeds = _relevant_feed_status(result, feed_status)
    snapshot = _feed_snapshot(relevant_feeds)
    inventory_time = inv.collection_time_utc or inv.collection_time_local or "Not provided"
    validity = (
        f"Findings reflect inventory collected at {inventory_time} and feed data checked at {snapshot['checked']}. "
        f"Revalidate after {snapshot['valid_until']} or immediately when newer vendor/security feeds are installed."
    )
    if snapshot["issues"]:
        validity += f" This report used {len(snapshot['issues'])} stale, missing, invalid, or failed feed record(s)."
    if result.assessment_state != "complete":
        validity = (
            f"No full vulnerability-assessment validity can be claimed. {result.assessment_reason} "
            f"Feed recheck due {snapshot['valid_until']} describes feed freshness only, not host security or coverage."
        )
    elif _coverage_issues(result):
        validity += " Coverage limitations: " + _coverage_text(result)
    unassessed = result.assessment_state == "incomplete" and not findings
    unavailable = "Not assessed"

    rows = [
        ("Report Generated", datetime.now().replace(microsecond=0)),
        ("Host", inv.hostname),
        ("Operating System", inv.os.pretty_name or inv.os.scope_key),
        ("Detected OS Scope", f"{result.detected_os_family}/{result.detected_os_product}"),
        ("Analysis OS Selection", f"{result.requested_os_family} → {inv.os.family}/{inv.os.product}"),
        ("Running Kernel", inv.running_kernel),
        ("Collected UTC", inv.collection_time_utc),
        ("Installed Packages", len(inv.packages)),
        ("Package Manager", ", ".join(managers) or "Unknown"),
        ("Standalone Software", len(inv.software)),
        ("Third-Party Packages", third_party_count),
        ("Package Origin Unknown", unknown_origin_count),
        ("Vulnerability Findings", unavailable if unassessed else len(confirmed_groups)),
        ("Scope / Evidence Review Candidates", len(review_groups)),
        ("Report Rows (Confirmed + Review)", len(finding_groups)),
        ("Unique CVEs", unavailable if unassessed else len(unique_cves)),
        ("Informational Inventory Records", len(inv.packages) + len(inv.software)),
        ("Critical", unavailable if unassessed else severity_counts["Critical"]),
        ("High", unavailable if unassessed else severity_counts["High"]),
        ("Medium", unavailable if unassessed else severity_counts["Medium"]),
        ("Low", unavailable if unassessed else severity_counts["Low"]),
        ("Known Exploited (CISA KEV)", sum(1 for items in confirmed_groups if any(x.kev for x in items))),
        ("Review Required", len(review_groups)),
        ("Feed Snapshot Checked UTC", snapshot["checked"]),
        ("Feed Valid Through UTC", snapshot["valid_until"]),
        ("Feed Freshness", snapshot["status"]),
        ("Finding Validity", validity),
        ("Analysis Status", _analysis_status(result, findings, relevant_feeds)),
        ("Metadata Collector", Path(inv.context_file).name if inv.context_file else "Embedded in input"),
        ("Latest Applicable Advisory", inv.assessment.get("latest_advisory", "Not recorded")),
        ("Packages in Applicable Advisory Sources", inv.assessment.get("evaluated_packages", "Not recorded")),
    ]

    ws.append(["Item", "Value"])
    for row in rows:
        ws.append(row)
    _style_header(ws)
    _finish_sheet(ws, {"A": 34, "B": 110})
    # Match the approved compact Summary, without changing the other sheets.
    for row in ws.iter_rows(min_row=2):
        lines = max(
            sum(max(1, math.ceil(len(line) / (32 if cell.column == 1 else 108))) for line in str(cell.value or "").split("\n"))
            for cell in row
        )
        ws.row_dimensions[row[0].row].height = min(409, max(15, lines * 15))
    ws.row_dimensions[28].height = max(60, ws.row_dimensions[28].height)
    ws["B2"].number_format = "mm-dd-yy"
    ws["B2"].alignment = Alignment(horizontal="left", vertical="top")
    # Make the key status row easier to spot.
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == "Analysis Status":
            ws.cell(row, 1).font = Font(bold=True)
            ws.cell(row, 2).font = Font(bold=True)
            ws.cell(row, 2).fill = PatternFill("solid", fgColor="FFF2CC")
    _add_table(ws, "SummaryTable")
    _summary_pie(ws, unassessed=unassessed, rated_total=sum(severity_counts[key] for key in ("Critical", "High", "Medium", "Low")), unrated=severity_counts["Info"])


def _summary_pie(ws, *, unassessed: bool, rated_total: int, unrated: int) -> None:
    """Native editable chart styled from the approved workbook; no sample data."""
    template = Path(__file__).with_name("assets") / "summary_pie.xml"
    chart = read_chart(ChartSpace.from_tree(fromstring(template.read_bytes())))
    chart.style = 2
    series = chart.series[0]
    series.cat.strRef.f = "'Summary'!$A$19:$A$22"
    series.val.numRef.f = "'Summary'!$B$19:$B$22"
    series.val.numRef.numCache = None
    if unassessed:
        chart.title = "Vulnerabilities — Not assessed"
    elif not rated_total:
        chart.title = "Vulnerabilities — No rated findings"
    elif unrated:
        chart.title = f"Vulnerabilities — rated findings ({unrated} unrated excluded)"
    chart.anchor = TwoCellAnchor(
        editAs="oneCell",
        _from=AnchorMarker(col=4, colOff=85725, row=2, rowOff=47625),
        to=AnchorMarker(col=11, colOff=158505, row=17, rowOff=73905),
    )
    ws.add_chart(chart)


def _vulnerability_headers() -> list[str]:
    return [
        "NUM", "CVE", "Severity", "Vendor Severity", "Finding Class", "Components", "Source Packages", "Matched Installed Packages",
        "Installed Version", "Fixed Version", "Advisory", "Status", "Confidence", "CVSS",
        "Description", "Solution", "Product Scope / CPE", "Evidence Source",
        "CISA KEV", "Review Required", "References",
    ]


def _vulnerability_row(index: int, items: list) -> list:
    raw: dict = {}
    for finding in items:
        for key, value in (finding.evidence.raw or {}).items():
            existing = raw.get(key)
            if isinstance(existing, list) and isinstance(value, list):
                raw[key] = list(dict.fromkeys(existing + value))
            elif isinstance(existing, dict) and isinstance(value, dict):
                raw[key] = {**existing, **value}
            elif existing is None:
                raw[key] = value
    matched = raw.get("matched_binary_packages") or list(dict.fromkeys(x.package_name for x in items))
    installed_versions = raw.get("installed_binary_versions") or {}
    fixed_versions = raw.get("fixed_binary_versions") or {}
    matched_text = "\n".join(
        f"{name}={installed_versions[name]}" if isinstance(installed_versions, dict) and name in installed_versions else str(name)
        for name in matched
    )
    scope_values = list(raw.get("affected_platforms") or []) + list(raw.get("affected_cpes") or [])
    if raw.get("scope_note"):
        scope_values.append(str(raw["scope_note"]))
    component_versions = list(dict.fromkeys((x.package_name, x.installed_evr) for x in items if x.installed_evr))
    installed_text = component_versions[0][1] if len(component_versions) == 1 else "\n".join(f"{name}={value}" for name, value in component_versions)
    component_fixed = list(dict.fromkeys((x.package_name, x.fixed_evr) for x in items if x.fixed_evr))
    fixed_text = component_fixed[0][1] if len(component_fixed) == 1 else "\n".join(f"{name}={value}" for name, value in component_fixed)
    if isinstance(fixed_versions, dict) and len(set(fixed_versions.values())) > 1:
        fixed_text = "\n".join(f"{name}={value}" for name, value in sorted(fixed_versions.items()))
    cves = list(dict.fromkeys(x.cve for x in items if x.cve))
    cvss = list(dict.fromkeys(f"{x.cve}: {x.cvss}" for x in items if x.cve and x.cvss))
    references = list(dict.fromkeys(ref for x in items for ref in x.references))
    components = list(dict.fromkeys(x.package_name for x in items if x.package_name))
    source_packages = list(dict.fromkeys(
        x.source_package or x.source_rpm or str((x.evidence.raw or {}).get("source_package") or "")
        for x in items
        if x.source_package or x.source_rpm or (x.evidence.raw or {}).get("source_package")
    ))
    advisories = list(dict.fromkeys(x.evidence.advisory for x in items if x.evidence.advisory))
    confidences = list(dict.fromkeys(x.confidence for x in items if x.confidence))
    evidence_sources = list(dict.fromkeys(x.evidence.source for x in items if x.evidence.source))
    description = max((x.description for x in items if x.description), key=len, default="")
    solutions = list(dict.fromkeys(x.solution.strip() for x in items if x.solution and x.solution.strip()))
    if solutions:
        solution = "\n".join(solutions)
    elif _group_requires_review(items):
        solution = "Validate vendor product, repository, release and package scope before remediation. Candidate fixed builds and matched packages are listed in this row."
    else:
        solution = f"Apply {', '.join(advisories) or 'the authoritative vendor remediation'} for the confirmed scope and update all matched packages."
    return [
        index,
        "\n".join(cves),
        _group_severity(items),
        _group_vendor_severity(items),
        _group_class(items),
        "\n".join(components),
        "\n".join(source_packages),
        matched_text,
        installed_text,
        fixed_text,
        "\n".join(advisories),
        _group_status(items),
        "; ".join(confidences),
        "\n".join(cvss),
        description,
        solution,
        "\n".join(dict.fromkeys(scope_values)),
        "\n".join(evidence_sources),
        "Yes" if any(x.kev for x in items) else "No",
        "Yes" if any(x.review_required for x in items) else "No",
        "\n".join(references),
    ]


def _vulnerabilities(ws, result: AnalysisResult) -> None:
    groups = _finding_groups(result)
    headers = _vulnerability_headers()
    ws.append(headers)
    for idx, items in enumerate(groups, 1):
        ws.append(_vulnerability_row(idx, items))
    _style_header(ws)

    severity_col = headers.index("Severity") + 1
    status_col = headers.index("Status") + 1
    review_col = headers.index("Review Required") + 1
    for row in range(2, ws.max_row + 1):
        sev = str(ws.cell(row, severity_col).value or "Info")
        fill, font = SEVERITY.get(sev, SEVERITY["Info"])
        ws.cell(row, severity_col).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row, severity_col).font = Font(color=font, bold=True)
        status = str(ws.cell(row, status_col).value or "")
        if status in STATUS_FILL:
            ws.cell(row, status_col).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
            ws.cell(row, status_col).font = Font(bold=True)
        if ws.cell(row, review_col).value == "Yes":
            ws.cell(row, review_col).fill = PatternFill("solid", fgColor="FFF2CC")
            ws.cell(row, review_col).font = Font(bold=True)
        ws.row_dimensions[row].height = 96

    _finish_sheet(ws, {
        "A": 8, "B": 20, "C": 12, "D": 15, "E": 24, "F": 32, "G": 40, "H": 48,
        "I": 30, "J": 40, "K": 24, "L": 26, "M": 25, "N": 16,
        "O": 58, "P": 58, "Q": 70, "R": 30, "S": 12, "T": 16, "U": 70,
    })
    _add_table(ws, "VulnerabilitiesTable")

def _inventory(ws, result: AnalysisResult) -> None:
    headers = [
        "NUM", "Type", "Name / Product", "Version", "Release / Build", "Architecture",
        "Vendor / Origin", "Source Package / Path", "Third Party", "Notes",
    ]
    ws.append(headers)
    num = 0

    for pkg in result.inventory.packages:
        num += 1
        manager = str(getattr(pkg, "package_manager", "rpm") or "rpm").lower()
        third = _third_party_state(result, pkg)
        if not pkg.metadata_complete:
            package_type = "Package list"
            version, release, vendor, source = pkg.installed_version, "", "Unknown", "Not supplied"
            third = None
            notes = "Name and version imported. Not assessed: source-package mapping, architecture, installation status and OS release require a matching full collector."
        elif manager in {"deb", "dpkg"}:
            package_type = "DEB"
            version = pkg.installed_version
            release = ""
            vendor = pkg.vendor or pkg.origin
            source = pkg.source_name
            if pkg.source_installed_version:
                source += f" (= {pkg.source_installed_version})"
            details = ["PackageManager=dpkg"]
            if pkg.status:
                details.append(f"Status={pkg.status}")
            if pkg.maintainer:
                details.append(f"Maintainer={pkg.maintainer}")
            notes = "; ".join(details)
            if result.inventory.assessment.get("mode") in {"advisories", "historical"}:
                matched = pkg.source_name in result.inventory.assessment.get("matched_sources", [])
                notes += "; " + ("Source checked against published advisories; open/untriaged issues not covered" if matched else "Not assessed: source is not listed in the release's loaded advisories")
        elif manager in {"apk", "pacman"}:
            package_type = manager.upper()
            version, release = pkg.installed_version, ""
            vendor, source = pkg.vendor or "Unknown", pkg.source_package or "Not supplied"
            notes = f"PackageManager={manager}; installed inventory only. This distribution has no dedicated VA adapter."
            if pkg.validation_metadata:
                notes += f"; PACMAN_VALIDATION_METADATA={pkg.validation_metadata} (database metadata only; not verified)"
        else:
            package_type = "RPM"
            version = pkg.version
            release = pkg.release
            vendor = pkg.vendor
            source = pkg.source_rpm
            notes = f"EVR={pkg.evr}; NEVRA={pkg.nevra}"
            if pkg.signature_keyid:
                notes += f"; RPM_SIGNATURE_METADATA Key ID={pkg.signature_keyid} (header metadata only; not verified)"
        ws.append([
            num, package_type, pkg.name, version, release, pkg.arch, vendor,
            source, "Yes" if third is True else "No" if third is False else "Unknown", notes,
        ])

    for sw in result.inventory.software:
        num += 1
        ws.append([
            num, "Standalone", sw.product, sw.version, sw.build, sw.architecture,
            sw.vendor, sw.path, "N/A", sw.evidence,
        ])

    _style_header(ws)
    third_col = headers.index("Third Party") + 1
    type_col = headers.index("Type") + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, third_col).value == "Yes":
            ws.cell(row, third_col).fill = PatternFill("solid", fgColor="FFF2CC")
        if ws.cell(row, type_col).value in {"Standalone", "DEB"}:
            ws.cell(row, type_col).fill = PatternFill("solid", fgColor="D9EAF7")
            ws.cell(row, type_col).font = Font(bold=True)

    _finish_sheet(ws, {
        "A": 8, "B": 14, "C": 32, "D": 22, "E": 28, "F": 14,
        "G": 36, "H": 58, "I": 14, "J": 65,
    })
    _add_table(ws, "InventoryTable")


def write_report(result: AnalysisResult, output: str | Path, feed_status: list | None = None) -> Path:
    """Write the simplified three-sheet client report AND full assessment template report."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _summary(summary, result, feed_status)

    _vulnerabilities(wb.create_sheet("Vulnerabilities"), result)
    _inventory(wb.create_sheet("Inventory"), result)

    # Never leave a half-written report (or destroy the previous report) when
    # saving is interrupted or the destination is open in Excel.
    fd, temporary = tempfile.mkstemp(prefix=".offlineva-", suffix=".xlsx", dir=output.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        os.replace(temporary, output)
    finally:
        wb.close()
        Path(temporary).unlink(missing_ok=True)

    # Full OS Security Assessment (same template used by enterprise reporting)
    try:
        from .full_report import write_full_assessment
        full_path = output.with_name(output.stem.replace("_cve_report", "") + "_OS_Security_Assessment.xlsx")
        if full_path.resolve() == output.resolve():
            full_path = output.with_name(output.stem + "_OS_Security_Assessment.xlsx")
        inv_text = Path(result.inventory.source_file) if result.inventory and result.inventory.source_file else None
        write_full_assessment(result, full_path, inventory_text=inv_text)
    except Exception as exc:
        # Full report is additive; never fail the primary CVE workbook.
        print(f"[WARN] full assessment report skipped: {exc}")

    return output
